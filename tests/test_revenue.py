from __future__ import annotations

import os
import json
import logging
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtCore import Qt
from PySide6.QtWidgets import QApplication, QDialog

from app.dialogs.revenue_dialog import RevenueDialog, RevenueEditDialog
from app.daily_import_ui import JsonExtractDialog
from app.config import AppConfig, get_default_settings
from app.database import Database
from app.main_window import MainWindow
from app.revenue import RevenueStore, parse_amount_vnd


def _row(sqt: str, container: str, amount, line: str, *, status: str = "OK") -> dict:
    return {
        "file_nguon": "Thang 10.xls",
        "ma_md5_file": "a" * 32,
        "loai_chung_tu": "Doanh thu",
        "trang_thai": status,
        "so_container": container,
        "canh_bao": [],
        "doanh_thu": {
            "sqt": sqt, "doi_tac": "Mía đường TN", "thanh_tien_gom_vat": amount,
            "sheet_nguon": "Mía đường TN", "dong_nguon": line,
        },
    }


class RevenueTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.store = RevenueStore(str(Path(self.temp.name) / "app_state.db"))
        self.store.init_db()

    def tearDown(self) -> None:
        self.temp.cleanup()

    def test_money_parses_vnd_without_float_rounding(self) -> None:
        self.assertEqual(parse_amount_vnd("75.239.928"), 75239928)
        self.assertEqual(parse_amount_vnd("75,239,928"), 75239928)
        self.assertEqual(parse_amount_vnd("75239928.0"), 75239928)
        self.assertIsNone(parse_amount_vnd("75,239.50"))

    def test_mixed_payload_keeps_incomplete_rows_and_reimport_keeps_edits(self) -> None:
        payload = {"du_lieu_boc_tach": [
            {"loai_chung_tu": "Bill", "so_container": "GAOU2133678"},
            _row("920", "VSGU2231272", 75239928, "9"),
            _row("921", "GAOU2133678", None, "10"),
        ]}
        self.assertEqual(self.store.ingest_payload(payload, "output.json"), 2)
        rows = self.store.list_rows()
        self.assertEqual({row["review_status"] for row in rows}, {"OK", "CAN_KIEM_TRA"})
        missing = next(row for row in rows if row["sqt"] == "921")
        self.store.save_row(
            missing["id"], partner="Mía đường TN", sqt="921",
            container="GAOU2133678", amount="77.497.560", warning_resolved=True,
            import_status="DA_NHAP",
        )
        self.assertEqual(self.store.ingest_payload(payload, "output.json"), 0)
        restored = next(row for row in self.store.list_rows() if row["sqt"] == "921")
        self.assertEqual(restored["amount_vnd"], 77497560)
        self.assertEqual(restored["review_status"], "OK")
        self.assertEqual(restored["import_status"], "DA_NHAP")
        with self.store._connect() as conn:
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM revenue_row_events").fetchone()[0], 1)

    def test_amended_source_flags_possible_duplicate(self) -> None:
        row = _row("920", "VSGU2231272", 75239928, "9")
        self.store.ingest_payload({"du_lieu_boc_tach": [row]}, "first.json")
        amended = dict(row, ma_md5_file="b" * 32)
        self.store.ingest_payload({"du_lieu_boc_tach": [amended]}, "second.json")
        rows = self.store.list_rows()
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]["review_status"], "OK")
        self.assertIn("Nghi trùng", rows[0]["warning_json"])

    def test_missing_partner_and_source_warning_do_not_block_revenue(self) -> None:
        row = _row("920", "VSGU2231272", 75239928, "9", status="Cần kiểm tra")
        row["doanh_thu"]["doi_tac"] = None
        row["canh_bao"] = ["Thiếu đối tác"]
        self.store.ingest_payload({"du_lieu_boc_tach": [row]}, "output.json")
        saved = self.store.list_rows()[0]
        self.assertEqual(saved["partner"], "")
        self.assertEqual(saved["review_status"], "OK")

    def test_existing_revenue_rows_recompute_review_status_on_upgrade(self) -> None:
        row = _row("920", "VSGU2231272", 75239928, "9")
        row["doanh_thu"]["doi_tac"] = None
        self.store.ingest_payload({"du_lieu_boc_tach": [row]}, "output.json")
        with self.store._connect() as conn:
            conn.execute("UPDATE revenue_rows SET review_status='CAN_KIEM_TRA'")
        self.store.init_db()
        self.assertEqual(self.store.list_rows()[0]["review_status"], "OK")

    def test_batch_save_is_atomic(self) -> None:
        self.store.ingest_payload({"du_lieu_boc_tach": [
            _row("920", "VSGU2231272", 75239928, "9"),
            _row("921", "GAOU2133678", 77497560, "10"),
        ]}, "output.json")
        rows = self.store.list_rows()
        changes = [{
            "id": row["id"], "partner": row["partner"], "sqt": row["sqt"],
            "container": row["container"], "amount": row["amount_vnd"],
            "import_status": "DA_NHAP",
        } for row in rows]
        changes[1]["import_status"] = "KHONG_HOP_LE"
        with self.assertRaises(ValueError):
            self.store.save_rows(changes)
        self.assertTrue(all(row["import_status"] == "CHUA_NHAP" for row in self.store.list_rows()))

    def test_dialog_stages_edits_and_bulk_status_until_one_save(self) -> None:
        app = QApplication.instance() or QApplication([])
        self.store.ingest_payload({"du_lieu_boc_tach": [
            _row("920", "VSGU2231272", 75239928, "9"),
            _row("920", "GAOU2133678", None, "10"),
        ]}, "output.json")
        dialog = RevenueDialog(self.store)
        self.assertEqual(dialog.table.rowCount(), 2)
        self.assertTrue(dialog.windowFlags() & Qt.WindowMaximizeButtonHint)
        self.assertEqual(dialog.table.item(0, 5).foreground().color().name(), "#dc2626")
        dialog.select_visible.click()
        self.assertEqual(len(dialog.selected_ids), 2)
        dialog.bulk_status.setCurrentIndex(dialog.bulk_status.findData("DA_NHAP"))
        dialog.bulk_apply.click()
        self.assertEqual(len(dialog.drafts), 2)
        self.assertTrue(all(row["import_status"] == "CHUA_NHAP" for row in self.store.list_rows()))
        missing_index = next(index for index, row_id in enumerate(dialog.visible_ids)
                             if dialog.rows[row_id]["container"] == "GAOU2133678")
        missing_id = dialog.visible_ids[missing_index]
        edited = {
            "partner": "", "sqt": "920", "container": "GAOU2133678",
            "amount_vnd": 77497560, "import_status": "DA_NHAP",
        }
        with patch.object(RevenueEditDialog, "exec", return_value=QDialog.Accepted), patch.object(
            RevenueEditDialog, "values", return_value=edited
        ):
            dialog._on_cell_clicked(missing_index, 4)
        self.assertEqual(dialog.rows[missing_id]["review_status"], "OK")
        self.assertIsNone(next(row for row in self.store.list_rows()
                               if row["id"] == missing_id)["amount_vnd"])
        dialog.save_button.click()
        rows = self.store.list_rows()
        changed = next(row for row in rows if row["id"] == missing_id)
        self.assertEqual(changed["amount_vnd"], 77497560)
        self.assertEqual(changed["review_status"], "OK")
        self.assertTrue(all(row["import_status"] == "DA_NHAP" for row in rows))
        self.assertEqual(dialog.table.item(0, 5).foreground().color().name(), "#15803d")
        with self.store._connect() as conn:
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM revenue_row_events").fetchone()[0], 2)
        dialog.close()
        del app

    def test_revenue_month_filter_uses_daily_sqt_and_clears_selection(self) -> None:
        app = QApplication.instance() or QApplication([])
        daily_file = Path(self.temp.name) / "daily.xlsx"
        workbook = Workbook()
        july = workbook.active
        july.title = "Tháng 7"
        july.append(["SQT PM"])
        july.append([920])
        august = workbook.create_sheet("Tháng 8")
        august.append(["SQT PM"])
        august.append([921])
        legacy = workbook.create_sheet("Thong_Tin_Quyet_Toan")
        legacy.append(["SQT PM"])
        legacy.append([999])
        workbook.save(daily_file)
        workbook.close()
        self.store.ingest_payload({"du_lieu_boc_tach": [
            _row("920", "VSGU2231272", 75_239_928, "9"),
            _row("921", "GAOU2133678", 77_497_560, "10"),
            _row("999", "VIMU2250703", 74_082_168, "11"),
        ]}, "output.json")

        dialog = RevenueDialog(self.store, daily_file=str(daily_file))
        self.assertEqual(dialog.month_filter.currentData(), "ALL")
        self.assertEqual(dialog.table.rowCount(), 3)
        dialog.month_filter.setCurrentIndex(dialog.month_filter.findData("Tháng 7"))
        self.assertEqual(dialog.table.rowCount(), 1)
        self.assertEqual(dialog.table.item(0, 1).text(), "920")
        dialog.select_visible.click()
        dialog.month_filter.setCurrentIndex(dialog.month_filter.findData("Tháng 8"))
        self.assertEqual(dialog.table.rowCount(), 1)
        self.assertEqual(dialog.table.item(0, 1).text(), "921")
        self.assertEqual(len(dialog.selected_ids), 0)
        dialog.select_visible.click()
        self.assertEqual(len(dialog.selected_ids), 1)
        dialog.month_filter.setCurrentIndex(dialog.month_filter.findData("UNKNOWN"))
        self.assertEqual(dialog.table.rowCount(), 1)
        self.assertEqual(dialog.table.item(0, 1).text(), "999")
        self.assertEqual(len(dialog.selected_ids), 0)
        self.assertIn("Đã chọn 0 dòng", dialog.summary.text())
        dialog.close()
        del app

    def test_revenue_status_and_search_filters_clear_selection(self) -> None:
        app = QApplication.instance() or QApplication([])
        self.store.ingest_payload({"du_lieu_boc_tach": [
            _row("920", "VSGU2231272", 75_239_928, "9"),
            _row("921", "GAOU2133678", 77_497_560, "10"),
        ]}, "output.json")
        dialog = RevenueDialog(self.store)
        dialog.select_visible.click()

        dialog.import_filter.setCurrentIndex(dialog.import_filter.findData("CHUA_NHAP"))
        self.assertEqual(dialog.selected_ids, set())
        self.assertTrue(all(dialog.table.item(row, 0).checkState() == Qt.Unchecked for row in range(2)))

        dialog.select_visible.click()
        dialog.search.setText("921")
        self.assertEqual(dialog.selected_ids, set())
        self.assertEqual(dialog.table.rowCount(), 1)
        self.assertIn("Đã chọn 0 dòng", dialog.summary.text())
        dialog.close()
        del app

    def test_revenue_sqt_sort_keeps_checkbox_actions_on_correct_row(self) -> None:
        app = QApplication.instance() or QApplication([])
        self.store.ingest_payload({"du_lieu_boc_tach": [
            _row("10", "VSGU2231272", 100, "9"),
            _row("2", "GAOU2133678", 200, "10"),
            _row("9", "VIMU2250703", 900, "11"),
        ]}, "output.json")
        dialog = RevenueDialog(self.store)
        self.assertEqual(dialog.sort_combo.currentData(), "sqt_asc")
        self.assertEqual([dialog.table.item(i, 1).text() for i in range(3)], ["2", "9", "10"])
        dialog.table.item(0, 0).setCheckState(Qt.Checked)
        selected_id = next(iter(dialog.selected_ids))
        self.assertEqual(dialog.rows[selected_id]["sqt"], "2")
        dialog.sort_combo.setCurrentIndex(dialog.sort_combo.findData("sqt_desc"))
        self.assertEqual([dialog.table.item(i, 1).text() for i in range(3)], ["10", "9", "2"])
        self.assertEqual(dialog.table.item(2, 0).checkState(), Qt.Checked)
        dialog.bulk_status.setCurrentIndex(dialog.bulk_status.findData("DA_NHAP"))
        dialog.bulk_apply.click()
        self.assertEqual(dialog.rows[selected_id]["import_status"], "DA_NHAP")
        self.assertEqual(len(dialog.drafts), 1)
        dialog.close()
        del app

    def test_editor_popup_only_stages_values(self) -> None:
        app = QApplication.instance() or QApplication([])
        self.store.ingest_payload({"du_lieu_boc_tach": [
            _row("920", "VSGU2231272", 75239928, "9"),
        ]}, "output.json")
        row = self.store.list_rows()[0]
        editor = RevenueEditDialog(row)
        editor.partner_edit.clear()
        editor.amount_edit.setText("75.239.928")
        self.assertEqual(editor.values()["amount_vnd"], 75239928)
        self.assertEqual(editor.values()["partner"], "")
        self.assertEqual(self.store.list_rows()[0]["partner"], "Mía đường TN")
        editor.close()
        del app

    def test_old_review_screen_preserves_but_does_not_show_revenue(self) -> None:
        app = QApplication.instance() or QApplication([])
        path = Path(self.temp.name) / "extract.json"
        payload = {"metadata": {}, "du_lieu_boc_tach": [
            {"loai_chung_tu": "Bill", "file_nguon": "bill.pdf", "so_container": "VSGU2231272"},
            _row("920", "VSGU2231272", 75239928, "9"),
        ]}
        path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
        dialog = JsonExtractDialog(str(path))
        self.assertEqual(len(dialog.rows), 1)
        self.assertEqual(dialog.summary_values["total"].text(), "1")
        dialog.save()
        saved = json.loads(path.read_text(encoding="utf-8"))
        self.assertEqual(len(saved["du_lieu_boc_tach"]), 2)
        self.assertEqual(saved["du_lieu_boc_tach"][1]["doanh_thu"]["sqt"], "920")
        del app

    def test_new_revenue_output_routes_to_step_four_without_old_review(self) -> None:
        app = QApplication.instance() or QApplication([])
        root = Path(self.temp.name)
        config = AppConfig(get_default_settings(str(root)), str(root / "Config" / "settings.json"))
        config.ensure_folders()
        database = Database(config.database_path)
        database.init_db()
        path = root / "output" / "boc_tach_doanh_thu.json"
        path.write_text(json.dumps({"metadata": {}, "du_lieu_boc_tach": [
            _row("920", "VSGU2231272", 75239928, "9"),
        ]}, ensure_ascii=False), encoding="utf-8")
        record_id = database.insert_processed_file(str(path), "READY_FOR_REVIEW")
        with patch("app.main_window.OutputWatcher.start"), patch("app.main_window.OutputWatcher.stop"), patch.object(
            MainWindow, "_start_excel_lock_check"
        ), patch.object(MainWindow, "_open_extract_editor") as old_review:
            window = MainWindow(config, database, logging.getLogger("test_revenue_ui"))
            window.on_output_ready({
                "id": record_id, "working_path": str(path), "file_name": path.name,
                "status": "READY_FOR_REVIEW",
            })
            self.assertEqual(len(window.revenue_store.list_rows()), 1)
            self.assertFalse(window.btn_view_extract.isEnabled())
            self.assertFalse(window.btn_import_daily.isEnabled())
            self.assertIn("1 dòng doanh thu", window.lbl_input_revenue_count.text())
            old_review.assert_not_called()
            window.close()
        del app


if __name__ == "__main__":
    unittest.main()
