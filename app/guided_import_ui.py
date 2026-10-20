"""Cửa sổ xử lý xung đột theo nhóm trước khi nhập file theo dõi."""

from __future__ import annotations

import copy
import html
from datetime import date, datetime
from typing import Any

from openpyxl import load_workbook
from PySide6.QtCore import QEvent, QTimer, Qt
from PySide6.QtGui import QColor, QKeySequence
from PySide6.QtWidgets import (
    QApplication, QComboBox, QHBoxLayout, QHeaderView, QLabel,
    QLineEdit, QMessageBox, QPushButton, QTableWidget, QTableWidgetItem,
    QTabWidget, QTextBrowser, QVBoxLayout, QWidget,
)

from . import file_utils
from .daily_import import (
    DOC_EXPENSE, EXPENSE_SHEET, DailyImportError, ImportAnalysis,
    ISSUE_FIELD_CONFLICT, ISSUE_MISSING_CLOSE_DATE, ISSUE_MULTIPLE_BILLS,
    ISSUE_MULTIPLE_TARGETS, ISSUE_NO_TARGET, RESOLUTION_FIELDS,
    RESOLUTION_SELECT, RESOLUTION_SKIP, ROW_DEFER_PREFIX, DailyImportService,
    FieldConflict, ImportIssue,
    IssueResolution, load_extract_payload, save_extract_payload, parse_date, parse_number,
)
from .dialogs.base_dialog import ResizableDialog
from .dialogs.notifications import notify


EDIT_KINDS = {ISSUE_MISSING_CLOSE_DATE, ISSUE_NO_TARGET}
CHOOSE_KINDS = {ISSUE_MULTIPLE_BILLS, ISSUE_MULTIPLE_TARGETS}
STYLE = """
QWidget { font-family: 'Segoe UI', Arial; font-size: 10pt; color: #1D304A; }
QDialog { background: #F4F7FB; }
QLabel#title { font-size: 17pt; font-weight: 700; color: #102541; }
QLabel#heading { font-size: 12pt; font-weight: 700; color: #102541; }
QLabel#hint { color: #5E718B; }
QLabel#summary { background: white; border: 1px solid #DCE5F0; border-radius: 7px; padding: 5px 9px; color: #315170; }
QLabel#notice { background: #EAF3FF; color: #245A9D; padding: 5px 8px; border-radius: 7px; }
QLabel#detail { background: #F6F9FD; border: 1px solid #DCE5F0; border-radius: 7px; padding: 6px 8px; }
QPushButton { background: white; border: 1px solid #C6D3E2; border-radius: 7px; padding: 5px 10px; min-height: 22px; }
QPushButton:hover { background: #EDF3FB; }
QPushButton:disabled { color: #98A7B8; background: #F2F4F7; }
QPushButton[primary='true'] { color: white; background: #285DB7; border-color: #285DB7; font-weight: 700; }
QPushButton[quiet='true'] { color: #526983; background: #F8FAFD; border-color: #D6E0EC; }
QLineEdit, QComboBox { background: white; border: 1px solid #C7D3E2; border-radius: 6px; padding: 5px; }
QTableWidget { background: white; alternate-background-color: #FAFCFF; gridline-color: #E4EAF2; selection-background-color: #DDEBFF; selection-color: #172C4A; }
QTableWidget::item:focus { outline: 2px solid #285DB7; }
QTextBrowser { background: #F8FAFE; border: 1px solid #DCE5F0; border-radius: 7px; padding: 10px; }
QHeaderView::section { background: #EDF2F8; padding: 8px; border: 0; border-right: 1px solid #DCE5F0; }
QTabWidget::pane { border: 1px solid #DCE5F0; background: white; }
QTabBar::tab { background: #E9EFF7; padding: 7px 14px; margin-right: 3px; }
QTabBar::tab:selected { background: white; color: #285DB7; font-weight: 700; }
QTableWidget#conflictTable { background: #FFFFFF; alternate-background-color: #FAFCFF; gridline-color: #DFE8F2; selection-background-color: #EDF4FF; font-size: 10pt; }
QTableWidget#conflictTable::item { padding: 4px 7px; }
QTableWidget#conflictTable QHeaderView::section { background: #EAF1F8; color: #355675; font-size: 10pt; font-weight: 400; border-right: 1px solid #DCE5EF; border-bottom: 1px solid #DCE5EF; padding: 5px; }
QComboBox#conflictRule, QLineEdit#conflictCustom { min-height: 29px; max-height: 29px; font-size: 10pt; }
"""


def label(value: str, role: str = "") -> QLabel:
    widget = QLabel(value)
    widget.setWordWrap(True)
    if role:
        widget.setObjectName(role)
    return widget


def button(value: str, action, primary: bool = False) -> QPushButton:
    widget = QPushButton(value)
    if primary:
        widget.setProperty("primary", "true")
    widget.clicked.connect(action)
    return widget


def read_only(value: Any) -> QTableWidgetItem:
    result = QTableWidgetItem(str(value) if value not in (None, "") else "—")
    result.setFlags(result.flags() & ~Qt.ItemIsEditable)
    return result


def shown(value: Any) -> str:
    if isinstance(value, (datetime, date)):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        number = float(value)
        if number.is_integer():
            return f"{int(number):,}".replace(",", ".")
        return f"{number:,.2f}".replace(",", "_").replace(".", ",").replace("_", ".").rstrip("0").rstrip(",")
    if isinstance(value, str):
        try:
            return datetime.strptime(value, "%Y-%m-%d").strftime("%d/%m/%Y")
        except ValueError:
            pass
    return str(value) if value not in (None, "") else ""


class PasteTable(QTableWidget):
    def keyPressEvent(self, event):  # noqa: N802 - Qt API
        if event.matches(QKeySequence.Paste):
            values = QApplication.clipboard().text().replace("\r\n", "\n").strip("\n")
            start_row, start_col = self.currentRow(), self.currentColumn()
            if start_row >= 0 and start_col in (3, 4):
                for row_offset, line in enumerate(values.split("\n")):
                    for col_offset, value in enumerate(line.split("\t")):
                        row, col = start_row + row_offset, start_col + col_offset
                        if row < self.rowCount() and col in (3, 4):
                            item = QTableWidgetItem(value.strip())
                            self.setItem(row, col, item)
                return
        super().keyPressEvent(event)


class LiveImportSession:
    """Working copy of one import batch; Excel is written only after confirmation."""

    def __init__(self, service: DailyImportService, analysis: ImportAnalysis):
        self.service = service
        self.analysis = analysis
        self.payload = load_extract_payload(analysis.output_path)
        self.resolutions = dict(analysis.resolutions)
        self.workbook_rows = self._read_workbook_rows(analysis.daily_path)

    @staticmethod
    def _read_workbook_rows(path: str) -> dict[tuple[str, int], dict[str, Any]]:
        result = {}
        if file_utils.is_file_locked(path):
            raise DailyImportError(
                "File Excel đang được sử dụng hoặc có file khóa cần kiểm tra."
            )
        try:
            book = load_workbook(path, read_only=True, data_only=True)
            try:
                for sheet in book.worksheets:
                    rows = sheet.iter_rows(values_only=True)
                    next(rows, ())
                    mapping = DailyImportService._header_map(sheet)
                    if "SQT PM" not in mapping or "Số Container" not in mapping:
                        continue
                    for row_number, values in enumerate(rows, start=2):
                        if any(value is not None for value in values):
                            result[(sheet.title, row_number)] = {
                                name: values[column - 1] if column <= len(values) else None
                                for name, column in mapping.items()
                            }
            finally:
                book.close()
        except (OSError, ValueError):
            return {}
        return result

    def reanalyze(self):
        self.analysis = self.service.analyze(
            self.analysis.output_path, self.analysis.daily_path,
            processed_file_id=self.analysis.processed_file_id,
            payload_override=self.payload, resolutions=self.resolutions,
        )
        self.resolutions = dict(self.analysis.resolutions)
        return self.analysis

def user_issue_name(issue: ImportIssue) -> str:
    if issue.kind == ISSUE_MULTIPLE_BILLS:
        return "A1 · Chọn Bill"
    if issue.kind == ISSUE_MULTIPLE_TARGETS:
        return "B2 · Chọn số quyết toán" if issue.doc_type == DOC_EXPENSE else "A2 · Chọn số quyết toán"
    if issue.kind == ISSUE_MISSING_CLOSE_DATE:
        return "A3 · Bổ sung ngày đóng"
    if issue.kind == ISSUE_NO_TARGET:
        return "B1 · Chưa tìm thấy số quyết toán"
    if issue.kind == ISSUE_FIELD_CONFLICT:
        return "B3 · Số liệu khác file theo dõi" if issue.doc_type == DOC_EXPENSE else "A4 · Số liệu khác file theo dõi"
    return issue.title


def user_candidate_name(candidate) -> str:
    if candidate.kind == "BILL":
        bill = candidate.values.get("Số B/L") or "chưa có số"
        vessel = candidate.values.get("Tên tàu") or "chưa có tên tàu"
        return f"Bill {bill} · {vessel} · dòng {candidate.json_index + 1}"
    return f"Số quyết toán {candidate.sqt} · {candidate.sheet_name}, dòng {candidate.row_number}"


DEFERRED_BG = QColor("#FFF1B8")


class GuidedImportResolutionDialog(ResizableDialog):
    def __init__(self, service: DailyImportService, analysis: ImportAnalysis, parent=None):
        super().__init__(parent)
        self.sandbox = LiveImportSession(service, analysis)
        self.final_analysis: ImportAnalysis | None = None
        self.saved_only = False
        self._has_saved_draft = False
        self._dirty = False
        self.setWindowTitle("Xử lý dữ liệu trước khi nhập file theo dõi")
        self.setStyleSheet(STYLE)
        self.resize(1390, 860)
        self.setMinimumSize(1100, 710)
        self.edit_drafts: dict[str, tuple[str, str]] = {}
        self.choice_drafts: dict[str, str] = {}
        self.field_drafts: dict[str, tuple[str, str]] = {}
        self.edit_ids: list[str] = []
        self.choice_ids: list[str] = []
        self.diff_rows: list[tuple[str, Any]] = []
        self.choice_boxes: list[QComboBox] = []
        self.diff_boxes: list[QComboBox] = []
        self.diff_inputs: list[QLineEdit] = []
        self._refreshing = False
        self._build()
        self.refresh_ui(capture=False)

    @property
    def analysis(self) -> ImportAnalysis:
        return self.sandbox.analysis

    def _build(self):
        outer = QVBoxLayout(self)
        outer.setContentsMargins(10, 7, 10, 7)
        outer.setSpacing(4)
        outer.addWidget(label("Xử lý dữ liệu trước khi nhập", "title"))
        outer.addWidget(label("Kiểm tra và chọn cách xử lý trước khi ghi file theo dõi. Bạn có thể lưu để tiếp tục sau.", "notice"))
        self.summary = label("", "summary")
        outer.addWidget(self.summary)
        self.tabs = QTabWidget()
        self.tabs.addTab(self._edit_page(), "1. Sửa thông tin")
        self.tabs.addTab(self._choice_page(), "2. Chọn Bill / số quyết toán")
        self.tabs.addTab(self._diff_page(), "3. Chọn số liệu")
        self.tabs.addTab(self._review_page(), "4. Xem trước")
        self.tabs.currentChanged.connect(self.on_tab_change)
        outer.addWidget(self.tabs, 1)
        footer = QHBoxLayout()
        self.feedback = label("Chọn một nhóm để bắt đầu. Các mục có thể chuyển nhóm sau khi PM kiểm tra lại.", "hint")
        footer.addWidget(self.feedback, 1)
        close_button = button("Đóng", self.reject)
        close_button.setProperty("quiet", "true")
        footer.addWidget(close_button)
        self.finish_button = button("Xác nhận nhập vào file theo dõi", self.finish, True)
        footer.addWidget(self.finish_button)
        outer.addLayout(footer)

    def _page(self, heading: str, hint: str):
        page = QWidget()
        layout = QVBoxLayout(page)
        layout.setContentsMargins(12, 9, 12, 8)
        layout.setSpacing(5)
        layout.addWidget(label(heading, "heading"))
        layout.addWidget(label(hint, "hint"))
        return page, layout

    def _edit_page(self):
        page, layout = self._page(
            "Sửa thông tin để PM tìm đúng dòng",
            "Sửa ngày và Container ngay trong bảng. Với Khoản chi, cột ‘Ngày PM đang dùng’ cho biết phải sửa Ngày chạy hay Ngày đóng. Có thể dán nhiều ô cùng lúc.",
        )
        self.edit_table = PasteTable(0, 6)
        self.edit_table.setHorizontalHeaderLabels(["Chứng từ", "Vấn đề", "Ngày PM đang dùng", "Ngày", "Container", "Kết quả"])
        self.edit_table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.edit_table.verticalHeader().setVisible(False)
        self.edit_table.verticalHeader().setDefaultSectionSize(36)
        self.edit_table.setAlternatingRowColors(True)
        self.edit_table.setSelectionBehavior(QTableWidget.SelectRows)
        self.edit_table.setSelectionMode(QTableWidget.ExtendedSelection)
        layout.addWidget(self.edit_table, 1)
        self.edit_detail, self.edit_preview = self._add_detail_panel(layout, "Dữ liệu liên quan trong file theo dõi")
        self.edit_table.itemSelectionChanged.connect(self.update_edit_context)
        self.edit_table.itemChanged.connect(self.update_edit_actions)
        actions = QHBoxLayout()
        self.edit_defer = button("Xử lý sau", lambda: self.defer_selected("edit"))
        self.edit_defer.setProperty("quiet", "true")
        actions.addWidget(self.edit_defer)
        self.edit_undo = button("Hoàn tác", lambda: self.undo_defer("edit"))
        actions.addWidget(self.edit_undo)
        actions.addStretch()
        self.edit_apply = button("Lưu dữ liệu", self.apply_edits, True)
        actions.addWidget(self.edit_apply)
        layout.addLayout(actions)
        return page

    def _choice_page(self):
        page, layout = self._page(
            "Chọn Bill hoặc dòng quyết toán cho từng chứng từ",
            "Mỗi dòng có các lựa chọn riêng. Sau khi chọn, PM sẽ ghép lại và kiểm tra tiếp xem số liệu có khác file theo dõi không.",
        )
        self.choice_table = QTableWidget(0, 6)
        self.choice_table.setHorizontalHeaderLabels(["Chứng từ", "Vấn đề", "Ngày", "Container", "Chọn kết quả phù hợp", "Kết quả"])
        self.choice_table.horizontalHeader().setSectionResizeMode(QHeaderView.Interactive)
        self.choice_table.horizontalHeader().setStretchLastSection(False)
        self.choice_table.horizontalHeader().setFixedHeight(36)
        self.choice_table.verticalHeader().setVisible(False)
        self.choice_table.verticalHeader().setDefaultSectionSize(42)
        self.choice_table.setAlternatingRowColors(True)
        self.choice_table.setSelectionBehavior(QTableWidget.SelectRows)
        self.choice_table.setSelectionMode(QTableWidget.ExtendedSelection)
        self.choice_table.setEditTriggers(QTableWidget.NoEditTriggers)
        self.choice_table.viewport().installEventFilter(self)
        layout.addWidget(self.choice_table)
        self.choice_detail, self.choice_preview = self._add_detail_panel(
            layout, "So sánh các phương án — bấm vào dòng để chọn", large=True,
        )
        self.choice_detail.setMaximumHeight(42)
        self.choice_preview.horizontalHeader().setFixedHeight(36)
        self.choice_preview.verticalHeader().setDefaultSectionSize(42)
        self.choice_table.itemSelectionChanged.connect(self.update_choice_context)
        self.choice_preview.cellClicked.connect(self.choose_from_preview)
        actions = QHBoxLayout()
        self.choice_defer = button("Xử lý sau", lambda: self.defer_selected("choice"))
        self.choice_defer.setProperty("quiet", "true")
        actions.addWidget(self.choice_defer)
        self.choice_undo = button("Hoàn tác", lambda: self.undo_defer("choice"))
        actions.addWidget(self.choice_undo)
        actions.addStretch()
        self.choice_apply = button("Lưu dữ liệu", self.apply_choices, True)
        actions.addWidget(self.choice_apply)
        layout.addLayout(actions)
        return page

    def _diff_page(self):
        page, layout = self._page(
            "Chọn số liệu sẽ dùng trong file theo dõi",
            "Xem số liệu đang có và số liệu mới cạnh nhau. Có thể đặt một cách xử lý cho các trường đang hiện, rồi sửa ngoại lệ ở từng dòng.",
        )
        toolbar = QHBoxLayout()
        toolbar.setSpacing(6)
        toolbar.addWidget(label("Xem trường:"))
        self.field_filter = QComboBox()
        self.field_filter.addItem("Tất cả")
        self.field_filter.currentIndexChanged.connect(self.filter_fields)
        toolbar.addWidget(self.field_filter)
        toolbar.addWidget(label("Cách xử lý cho các dòng đang hiện:"))
        self.rule = QComboBox()
        self.rule.addItem("Chọn cách xử lý", "")
        self.rule.addItem("Giữ số liệu đang có", "CURRENT")
        self.rule.addItem("Dùng số liệu mới", "NEW")
        self.rule.currentIndexChanged.connect(self.update_rule_action)
        toolbar.addWidget(self.rule)
        self.rule_fill = button("Điền cho các dòng đang hiện", self.fill_rule)
        toolbar.addWidget(self.rule_fill)
        toolbar.addStretch()
        layout.addLayout(toolbar)
        self.diff_table = QTableWidget(0, 8)
        self.diff_table.setObjectName("conflictTable")
        self.diff_table.setHorizontalHeaderLabels(["Chứng từ", "SQT", "Trường", "File theo dõi", "Chứng từ mới", "Cách xử lý", "Giá trị nhập", "Kết quả"])
        self.diff_table.horizontalHeader().setSectionResizeMode(QHeaderView.Interactive)
        self.diff_table.horizontalHeader().setStretchLastSection(False)
        self.diff_table.horizontalHeader().setFixedHeight(36)
        self.diff_table.verticalHeader().setVisible(False)
        self.diff_table.verticalHeader().setDefaultSectionSize(42)
        self.diff_table.verticalHeader().setMinimumSectionSize(42)
        self.diff_table.setAlternatingRowColors(True)
        self.diff_table.setWordWrap(True)
        self.diff_table.setSelectionBehavior(QTableWidget.SelectRows)
        self.diff_table.setSelectionMode(QTableWidget.ExtendedSelection)
        self.diff_table.setEditTriggers(QTableWidget.NoEditTriggers)
        self.diff_table.viewport().installEventFilter(self)
        layout.addWidget(self.diff_table, 1)
        self.diff_table.itemSelectionChanged.connect(self.update_diff_context)
        detail_panel = QWidget()
        detail_layout = QVBoxLayout(detail_panel)
        detail_layout.setContentsMargins(0, 0, 0, 0)
        detail_layout.setSpacing(4)
        detail_layout.addWidget(label("Dòng đang có trong file theo dõi", "heading"))
        self.diff_detail = label("Chọn một dòng ở bảng trên để xem thông tin đối chiếu.", "detail")
        self.diff_detail.setMaximumHeight(68)
        detail_layout.addWidget(self.diff_detail)
        self.diff_preview = QTableWidget(0, 0)
        self.diff_preview.setMinimumHeight(78)
        self.diff_preview.verticalHeader().setVisible(False)
        self.diff_preview.verticalHeader().setDefaultSectionSize(42)
        self.diff_preview.horizontalHeader().setFixedHeight(36)
        self.diff_preview.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.diff_preview.setAlternatingRowColors(True)
        self.diff_preview.setEditTriggers(QTableWidget.NoEditTriggers)
        self.diff_preview.setSelectionBehavior(QTableWidget.SelectRows)
        detail_layout.addWidget(self.diff_preview, 1)
        self.diff_detail_panel = detail_panel
        detail_panel.setFixedHeight(185)
        layout.addWidget(detail_panel)
        actions = QHBoxLayout()
        self.scope = label("", "hint")
        actions.addWidget(self.scope)
        self.diff_defer = button("Xử lý sau", lambda: self.defer_selected("diff"))
        self.diff_defer.setProperty("quiet", "true")
        actions.addWidget(self.diff_defer)
        self.diff_undo = button("Hoàn tác", lambda: self.undo_defer("diff"))
        actions.addWidget(self.diff_undo)
        actions.addStretch()
        self.field_apply = button("Lưu dữ liệu", self.apply_fields, True)
        actions.addWidget(self.field_apply)
        layout.addLayout(actions)
        return page

    def _size_diff_columns(self):
        """Keep every conflict column visible while the table fills the dialog."""
        if not hasattr(self, "diff_table"):
            return
        minimums = (170, 58, 105, 120, 120, 160, 110, 80)
        weights = (18, 6, 11, 14, 14, 18, 11, 8)
        available = max(self.diff_table.viewport().width() - 2, sum(minimums))
        extra = available - sum(minimums)
        widths = [base + extra * weight // 100 for base, weight in zip(minimums, weights)]
        widths[-1] += available - sum(widths)
        for column, width in enumerate(widths):
            self.diff_table.setColumnWidth(column, width)

    def _size_choice_columns(self):
        if not hasattr(self, "choice_table"):
            return
        minimums = (160, 150, 95, 125, 250, 90)
        weights = (21, 20, 8, 12, 29, 10)
        available = max(self.choice_table.viewport().width() - 2, sum(minimums))
        extra = available - sum(minimums)
        widths = [base + extra * weight // 100 for base, weight in zip(minimums, weights)]
        widths[-1] += available - sum(widths)
        for column, width in enumerate(widths):
            self.choice_table.setColumnWidth(column, width)

    def _fit_choice_table(self):
        # The issue list grows with its rows; the candidate comparison uses the rest.
        rows = self.choice_table.rowCount()
        self.choice_table.setFixedHeight(max(82, min(36 + rows * 42 + 4, 230)))

    @staticmethod
    def _centered_table_control(widget: QWidget) -> QWidget:
        host = QWidget()
        cell = QHBoxLayout(host)
        cell.setContentsMargins(5, 4, 5, 4)
        cell.addWidget(widget)
        widget.setFixedHeight(30)
        return host

    def resizeEvent(self, event):  # noqa: N802 - Qt API
        super().resizeEvent(event)
        if hasattr(self, "diff_detail_panel"):
            self.diff_detail_panel.setFixedHeight(min(185, max(155, self.height() - 675)))
        self._size_choice_columns()
        self._size_diff_columns()

    def eventFilter(self, watched, event):  # noqa: N802 - Qt API
        if event.type() == QEvent.Resize:
            if hasattr(self, "choice_table") and watched is self.choice_table.viewport():
                self._size_choice_columns()
            elif hasattr(self, "diff_table") and watched is self.diff_table.viewport():
                self._size_diff_columns()
        return super().eventFilter(watched, event)

    def _add_detail_panel(self, layout: QVBoxLayout, heading: str, *, large: bool = False) -> tuple[QLabel, QTableWidget]:
        layout.addWidget(label(heading, "heading"))
        details = label("Chọn một dòng trong bảng phía trên để xem dữ liệu.", "detail")
        layout.addWidget(details)
        preview = QTableWidget(0, 0)
        preview.setMinimumHeight(150)
        if not large:
            preview.setMaximumHeight(225)
        preview.verticalHeader().setVisible(False)
        preview.verticalHeader().setDefaultSectionSize(35)
        preview.setAlternatingRowColors(True)
        preview.setEditTriggers(QTableWidget.NoEditTriggers)
        preview.setSelectionBehavior(QTableWidget.SelectRows)
        preview.horizontalHeader().setSectionResizeMode(QHeaderView.Interactive)
        layout.addWidget(preview, 1 if large else 0)
        return details, preview

    def _review_page(self):
        page, layout = self._page(
            "Xem trước kết quả nhập",
            "Số dòng dưới đây do PM tính từ các lựa chọn hiện tại. File theo dõi chỉ được ghi sau khi bạn xác nhận.",
        )
        self.review_summary = label("", "notice")
        layout.addWidget(self.review_summary)
        self.review_attention = label("", "detail")
        layout.addWidget(self.review_attention)
        self.review_table = QTableWidget(0, 6)
        self.review_table.setHorizontalHeaderLabels(["Phần dữ liệu", "Việc sẽ làm", "SQT", "Container", "Vị trí trong file", "Nội dung thay đổi"])
        self.review_table.horizontalHeader().setSectionResizeMode(QHeaderView.Interactive)
        self.review_table.horizontalHeader().setStretchLastSection(True)
        for col, width in enumerate((130, 125, 75, 145, 180)):
            self.review_table.setColumnWidth(col, width)
        self.review_table.setEditTriggers(QTableWidget.NoEditTriggers)
        self.review_table.setSelectionBehavior(QTableWidget.SelectRows)
        layout.addWidget(self.review_table, 1)
        layout.addWidget(label("Chi tiết dòng được chọn", "heading"))
        self.review_details = QTextBrowser()
        self.review_details.setMaximumHeight(185)
        layout.addWidget(self.review_details)
        self.review_table.itemSelectionChanged.connect(self.update_review_detail)
        return page

    def current_issues(self, kinds: set[str]) -> list[ImportIssue]:
        return [issue for issue in self.sandbox.analysis.issues if issue.kind in kinds]

    def source_row(self, issue: ImportIssue) -> dict[str, Any]:
        rows = self.sandbox.payload.get("du_lieu_boc_tach") or []
        if issue.json_indexes and 0 <= issue.json_indexes[0] < len(rows):
            row = rows[issue.json_indexes[0]]
            return row if isinstance(row, dict) else {}
        return {}

    def source_label(self, issue: ImportIssue) -> str:
        row = self.source_row(issue)
        reference = str(issue.reference or row.get("file_nguon") or "").strip()
        doc_label = issue.doc_label
        if reference.casefold().startswith(doc_label.casefold() + " "):
            return reference
        return f"{doc_label} {reference}".strip()

    def _issue_sqt(self, issue: ImportIssue, conflict: FieldConflict | None = None) -> str:
        if conflict and not conflict.target_row:
            return "Sẽ tạo"
        if conflict and conflict.target_row:
            values = self.sandbox.workbook_rows.get((conflict.target_sheet, conflict.target_row), {})
            return shown(values.get("SQT PM")) or "—"
        selected = self.sandbox.analysis.resolutions.get(issue.issue_id)
        if selected and selected.action == RESOLUTION_SELECT:
            candidate = next((item for item in issue.candidates if item.candidate_id == selected.candidate_id), None)
            if candidate and candidate.sqt:
                return shown(candidate.sqt)
        return "Chưa chọn" if issue.kind == ISSUE_MULTIPLE_TARGETS else "—"

    def _is_deferred(self, issue_id: str) -> bool:
        resolution = self.sandbox.analysis.resolutions.get(issue_id)
        return bool(resolution and resolution.action == RESOLUTION_SKIP)

    @staticmethod
    def _selected_rows(table: QTableWidget) -> list[int]:
        return sorted({index.row() for index in table.selectionModel().selectedRows()})

    @staticmethod
    def _paint_deferred(table: QTableWidget, row: int, deferred: bool) -> None:
        for column in range(table.columnCount()):
            item = table.item(row, column)
            if item:
                item.setBackground(DEFERRED_BG if deferred else QColor("#FFFFFF"))

    def effective_day_key(self, issue: ImportIssue) -> tuple[str, str]:
        row = self.source_row(issue)
        if issue.doc_type == DOC_EXPENSE and parse_date(row.get("ngay_chay")):
            return "ngay_chay", "Ngày chạy"
        return "ngay_dong", "Ngày đóng"

    def _source_summary(self, issue: ImportIssue) -> str:
        row = self.source_row(issue)
        fields = (
            [("Số hóa đơn", "so_hd"), ("Container", "so_container"),
             ("Ngày chạy", "ngay_chay"), ("Ngày đóng", "ngay_dong"),
             ("Đơn giá", "don_gia"), ("Nơi đóng", "noi_dong")]
            if issue.doc_type == DOC_EXPENSE else
            [("Số Bill", "so_bill"), ("Container", "so_container"),
             ("Tên tàu", "ten_tau"), ("Ngày chạy", "ngay_chay"),
             ("VT biển", "vt_bien"), ("Số Seal", "so_chi_seal")]
            if issue.doc_label == "Bill" else
            [("Container", "so_container"), ("Ngày đóng", "ngay_dong"),
             ("Số Seal", "so_chi_seal"), ("Số tấn", "so_tan"),
             ("Nơi đóng", "noi_dong"), ("Loại hàng", "loai_hang")]
        )
        cells = [
            f"<td width='33%'><b>{html.escape(name)}:</b> {html.escape(shown(row.get(key)) or 'Chưa có')}</td>"
            for name, key in fields
        ]
        grid = "".join("<tr>" + "".join(cells[index:index + 3]) + "</tr>" for index in range(0, len(cells), 3))
        file_name = html.escape(str(row.get("file_nguon") or "Chưa có"))
        related = (f"<br><b>Nhóm ghép:</b> {len(issue.json_indexes)} dòng chứng từ liên quan; "
                   "Xử lý sau sẽ áp dụng cho cả nhóm." if len(issue.json_indexes) > 1 else "")
        return (f"<b>Chứng từ đang xử lý · {html.escape(issue.doc_label)}</b>"
                f"<table width='100%' cellpadding='3' cellspacing='3'>{grid}</table>"
                f"<b>File:</b> {file_name}{related}")

    @staticmethod
    def _set_preview(table: QTableWidget, headers: list[str], rows: list[list[Any]], chosen: int = -1):
        table.setUpdatesEnabled(False)
        table.clear()
        table.setColumnCount(len(headers))
        table.setHorizontalHeaderLabels(headers)
        table.setRowCount(len(rows))
        for row_index, values in enumerate(rows):
            for col_index, value in enumerate(values):
                cell = read_only(shown(value))
                cell.setToolTip(str(value) if value not in (None, "") else "")
                if row_index == chosen:
                    cell.setBackground(QColor("#E5F1FF"))
                table.setItem(row_index, col_index, cell)
        table.verticalHeader().setDefaultSectionSize(46)
        table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeToContents)
        table.horizontalHeader().setStretchLastSection(True)
        table.setUpdatesEnabled(True)

    def _matching_workbook_rows(self, container: str, *, expense: bool = False):
        for (sheet, row_number), values in self.sandbox.workbook_rows.items():
            if (sheet == EXPENSE_SHEET) != expense:
                continue
            if str(values.get("Số Container") or "").strip().upper() == container.strip().upper():
                yield sheet, row_number, values

    def update_edit_context(self):
        if self._refreshing:
            return
        selected_ids = [self.edit_ids[row] for row in self._selected_rows(self.edit_table) if row < len(self.edit_ids)]
        self.edit_defer.setEnabled(any(not self._is_deferred(issue_id) for issue_id in selected_ids))
        self.edit_undo.setEnabled(any(self._is_deferred(issue_id) for issue_id in selected_ids))
        index = self.edit_table.currentRow()
        if not 0 <= index < len(self.edit_ids):
            self.edit_detail.setText("Chọn một dòng cần sửa để xem chứng từ và các dòng cùng Container.")
            self._set_preview(self.edit_preview, ["Dòng file theo dõi"], [])
            return
        issue = next(item for item in self.sandbox.analysis.issues if item.issue_id == self.edit_ids[index])
        source = self.source_row(issue)
        key, key_label = self.effective_day_key(issue)
        edited_day = self.edit_table.item(index, 3).text().strip()
        edited_container = self.edit_table.item(index, 4).text().strip().upper()
        draft = ""
        if (edited_day, edited_container) != (shown(source.get(key)), str(source.get("so_container") or "").upper()):
            draft = (f"<br><b>Đang sửa, chưa kiểm tra:</b> {html.escape(key_label)} → "
                     f"{html.escape(edited_day or 'Chưa có')} · Container → {html.escape(edited_container or 'Chưa có')}")
        self.edit_detail.setText(self._source_summary(issue) + draft +
                                 "<br><b>Gợi ý đối chiếu:</b> Các dòng dưới đây cùng Container; hãy kiểm tra ngày trước khi sửa.")
        rows = [
            [values.get("SQT PM"), values.get("Ngày Đóng"), values.get("Số Container"),
             values.get("Nơi đóng"), values.get("Loại hàng"), values.get("Người nhận"), f"{sheet}, dòng {number}"]
            for sheet, number, values in self._matching_workbook_rows(issue.container)
        ]
        self._set_preview(self.edit_preview,
                          ["Số quyết toán", "Ngày đóng", "Container", "Nơi đóng", "Loại hàng", "Người nhận", "Vị trí trong file"], rows)
        if not rows:
            self.edit_detail.setText(self.edit_detail.text() + "<br>Chưa có dòng cùng Container trong file theo dõi.")

    def update_choice_context(self):
        if self._refreshing:
            return
        selected_ids = [self.choice_ids[row] for row in self._selected_rows(self.choice_table) if row < len(self.choice_ids)]
        self.choice_defer.setEnabled(any(not self._is_deferred(issue_id) for issue_id in selected_ids))
        self.choice_undo.setEnabled(any(self._is_deferred(issue_id) for issue_id in selected_ids))
        index = self.choice_table.currentRow()
        if not 0 <= index < len(self.choice_ids):
            self.choice_detail.setText("Chọn một chứng từ phía trên để so sánh các phương án.")
            self._set_preview(self.choice_preview, ["Phương án"], [])
            return
        issue = next(item for item in self.sandbox.analysis.issues if item.issue_id == self.choice_ids[index])
        source = self.source_row(issue)
        self.choice_detail.setText(
            f"<b>{html.escape(self.source_label(issue))}</b>"
            f"　·　<b>File:</b> {html.escape(str(source.get('file_nguon') or 'Chưa có'))}"
            f"　·　<b>Container:</b> {html.escape(issue.container or 'Chưa có')}"
            f"　·　<b>Ngày:</b> {html.escape(shown(issue.match_date) or 'Chưa có')}"
        )
        selected_id = str(self.choice_boxes[index].currentData() or "")
        if issue.kind == ISSUE_MULTIPLE_BILLS:
            headers = ["Phương án", "Số Bill", "Ngày chạy", "Tên tàu", "VT biển", "Container", "File Bill"]
            rows = [
                [f"Bill {number}", candidate.values.get("Số B/L"), candidate.values.get("Ngày chạy"),
                 candidate.values.get("Tên tàu"), candidate.values.get("VT biển"),
                 candidate.values.get("Container"), candidate.values.get("File nguồn")]
                for number, candidate in enumerate(issue.candidates, 1)
            ]
        else:
            headers = ["Phương án", "Số quyết toán", "Ngày đóng", "Container", "Nơi đóng", "Loại hàng", "Người nhận", "Vị trí trong file"]
            rows = []
            for number, candidate in enumerate(issue.candidates, 1):
                values = self.sandbox.workbook_rows.get((candidate.sheet_name, candidate.row_number), candidate.values)
                rows.append([f"Dòng {number}", candidate.sqt, values.get("Ngày Đóng") or candidate.values.get("Ngày đóng"),
                             values.get("Số Container") or candidate.values.get("Container"), values.get("Nơi đóng"),
                             values.get("Loại hàng"), values.get("Người nhận"),
                             f"{candidate.sheet_name}, dòng {candidate.row_number}"])
        chosen = next((i for i, item in enumerate(issue.candidates) if item.candidate_id == selected_id), -1)
        self._set_preview(self.choice_preview, headers, rows, chosen)
        self.choice_preview.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.choice_preview.verticalHeader().setDefaultSectionSize(42)

    def choose_from_preview(self, row: int, _column: int):
        issue_row = self.choice_table.currentRow()
        if 0 <= issue_row < len(self.choice_boxes) and 0 <= row + 1 < self.choice_boxes[issue_row].count():
            self.choice_boxes[issue_row].setCurrentIndex(row + 1)
            self.update_choice_context()

    def update_diff_context(self):
        if self._refreshing:
            return
        selected_ids = [self.diff_rows[row][0] for row in self._selected_rows(self.diff_table) if row < len(self.diff_rows)]
        self.diff_defer.setEnabled(any(not self._is_deferred(issue_id) for issue_id in selected_ids))
        self.diff_undo.setEnabled(any(self._is_deferred(issue_id) for issue_id in selected_ids))
        index = self.diff_table.currentRow()
        if not 0 <= index < len(self.diff_rows):
            self.diff_detail.setText("Chọn một dòng ở bảng trên để xem thông tin đối chiếu.")
            self._set_preview(self.diff_preview, ["Dòng file theo dõi"], [])
            return
        issue_id, conflict = self.diff_rows[index]
        issue = next(item for item in self.sandbox.analysis.issues if item.issue_id == issue_id)
        source = self.source_row(issue)
        day_key, day_label = self.effective_day_key(issue)
        self.diff_detail.setText(
            f"Chứng từ đang xử lý: {issue.doc_label} · File: {source.get('file_nguon') or 'Chưa có'} · "
            f"{day_label}: {shown(source.get(day_key)) or 'Chưa có'} · Container: {issue.container or 'Chưa có'} · "
            f"Số Seal: {shown(source.get('so_chi_seal')) or 'Chưa có'} · "
            f"Số tấn: {shown(source.get('so_tan')) or 'Chưa có'} · "
            f"Nơi đóng: {shown(source.get('noi_dong')) or 'Chưa có'}\n"
            f"Đang so sánh: {conflict.field_name} · File theo dõi: {shown(conflict.current_value) or 'Chưa có'} · "
            f"Chứng từ mới: {shown(conflict.new_value) or 'Chưa có'}"
        )
        self.diff_detail.setToolTip(self.diff_detail.text())
        values = self.sandbox.workbook_rows.get((conflict.target_sheet, conflict.target_row), {})
        if issue.doc_type == DOC_EXPENSE:
            headers = ["Số quyết toán", "Ngày tháng", "Container", "Số hóa đơn", "Nơi đóng", "Đơn giá", "Tổng tiền", "Vị trí trong file"]
            row = [values.get("SQT PM"), values.get("Ngày tháng"), values.get("Số Container"),
                   values.get("Số HĐ"), values.get("Nơi đóng"), values.get("Đơn giá"),
                   values.get("Tổng tiền"), f"{conflict.target_sheet}, dòng {conflict.target_row}"]
        else:
            headers = ["Số quyết toán", "Ngày đóng", "Container", "Số Bill", "Số Seal", "Số tấn", "Nơi đóng", "Vị trí trong file"]
            row = [values.get("SQT PM"), values.get("Ngày Đóng"), values.get("Số Container"),
                   values.get("Số B/L"), values.get("Số chì/Seal"), values.get("Số tấn"),
                   values.get("Nơi đóng"), f"{conflict.target_sheet}, dòng {conflict.target_row}"]
        self._set_preview(self.diff_preview, headers, [row] if values else [])
        self.diff_preview.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.diff_preview.verticalHeader().setDefaultSectionSize(42)

    @staticmethod
    def _set_status(table: QTableWidget, row: int, column: int, value: str):
        item = table.item(row, column)
        if item and item.text() != value:
            blocked = table.blockSignals(True)
            item.setText(value)
            table.blockSignals(blocked)

    def update_edit_actions(self, *_args):
        if self._refreshing:
            return
        changed = 0
        for index, issue_id in enumerate(self.edit_ids):
            issue = next(item for item in self.sandbox.analysis.issues if item.issue_id == issue_id)
            source = self.source_row(issue)
            key, _ = self.effective_day_key(issue)
            day_item, container_item = self.edit_table.item(index, 3), self.edit_table.item(index, 4)
            if day_item is None or container_item is None:
                continue
            differs = (day_item.text().strip(), container_item.text().strip().upper()) != (
                shown(source.get(key)), str(source.get("so_container") or "").upper())
            changed += differs and not self._is_deferred(issue_id)
            self._set_status(self.edit_table, index, 5, "Xử lý sau" if self._is_deferred(issue_id) else "Chưa lưu" if differs else "Cần sửa")
        self.edit_apply.setEnabled(changed > 0)
        self.edit_apply.setText("Lưu dữ liệu")
        self.update_edit_context()

    def update_choice_actions(self):
        if self._refreshing:
            return
        changed = 0
        for index, (issue_id, box) in enumerate(zip(self.choice_ids, self.choice_boxes)):
            candidate_id = str(box.currentData() or "")
            saved = self.sandbox.analysis.resolutions.get(issue_id)
            is_saved = bool(saved and saved.action == RESOLUTION_SELECT and saved.candidate_id == candidate_id)
            pending = bool(candidate_id and not is_saved and not self._is_deferred(issue_id))
            changed += pending
            status = ("Xử lý sau" if saved and saved.action == RESOLUTION_SKIP else
                      "Chưa lưu" if pending else "Đã lưu" if is_saved else "Chưa chọn")
            self._set_status(self.choice_table, index, 5, status)
        self.choice_apply.setEnabled(changed > 0)
        self.choice_apply.setText("Lưu dữ liệu")

    def _current_field_choices(self, issue: ImportIssue) -> dict[str, dict[str, str]]:
        choices = {}
        for row, (issue_id, conflict) in enumerate(self.diff_rows):
            if issue_id != issue.issue_id:
                continue
            mode = str(self.diff_boxes[row].currentData() or "")
            value = self.diff_inputs[row].text().strip()
            if mode and (mode != "CUSTOM" or self._valid_custom(conflict, value)):
                choices[conflict.conflict_id] = {"mode": mode, "value": value}
        return choices

    def update_field_actions(self, *_args):
        if self._refreshing:
            return
        ready = 0
        for issue in self.current_issues({ISSUE_FIELD_CONFLICT}):
            saved = self.sandbox.analysis.resolutions.get(issue.issue_id)
            if saved and saved.action == RESOLUTION_FIELDS:
                for row, (issue_id, conflict) in enumerate(self.diff_rows):
                    if issue_id == issue.issue_id and not self.diff_boxes[row].currentData():
                        prior_mode = saved.field_choices.get(conflict.conflict_id, {}).get("mode")
                        prior_index = self.diff_boxes[row].findData(prior_mode)
                        if prior_index > 0:
                            blocked = self.diff_boxes[row].blockSignals(True)
                            self.diff_boxes[row].setCurrentIndex(prior_index)
                            self.diff_boxes[row].blockSignals(blocked)
                            self.diff_inputs[row].setEnabled(prior_mode == "CUSTOM")
            choices = self._current_field_choices(issue)
            complete = len(choices) == len(issue.conflicts)
            is_saved = bool(saved and saved.action == RESOLUTION_FIELDS and saved.field_choices == choices)
            ready += bool(complete and not is_saved and not (saved and saved.action == RESOLUTION_SKIP))
            for row, (issue_id, conflict) in enumerate(self.diff_rows):
                if issue_id != issue.issue_id:
                    continue
                mode = str(self.diff_boxes[row].currentData() or "")
                status = ("Chưa lưu" if complete and not is_saved else
                          "Đã lưu" if is_saved else
                          "Cần nhập giá trị" if mode == "CUSTOM" and conflict.conflict_id not in choices else
                          "Để sau" if saved and saved.action == RESOLUTION_SKIP and not mode else
                          "Chưa chọn")
                self._set_status(self.diff_table, row, 7, "Xử lý sau" if saved and saved.action == RESOLUTION_SKIP else status)
        self.field_apply.setEnabled(ready > 0)
        self.field_apply.setText("Lưu dữ liệu")
        self.update_rule_action()

    def update_rule_action(self, *_args):
        if self._refreshing:
            return
        fillable = sum(
            not self.diff_table.isRowHidden(row) and not box.currentData()
            and not (
                (saved := self.analysis.resolutions.get(self.diff_rows[row][0]))
                and saved.action == RESOLUTION_SKIP
            )
            for row, box in enumerate(self.diff_boxes)
        )
        self.rule_fill.setEnabled(bool(self.rule.currentData()) and fillable > 0)
        self.rule_fill.setText(f"Điền cho {fillable} trường đang hiện" if fillable else "Điền cho các dòng đang hiện")

    def _advance_if_done(self, step: int):
        unresolved = self.sandbox.analysis.unresolved_issues
        for next_step, kinds in ((0, EDIT_KINDS), (1, CHOOSE_KINDS), (2, {ISSUE_FIELD_CONFLICT})):
            if next_step > step and any(issue.kind in kinds for issue in unresolved):
                self.tabs.setCurrentIndex(next_step)
                return
        if not unresolved:
            self.tabs.setCurrentIndex(3)

    def _unapplied_step(self) -> int | None:
        if self.edit_apply.isEnabled():
            return 0
        if self.choice_apply.isEnabled():
            return 1
        for row, (issue_id, conflict) in enumerate(self.diff_rows):
            saved = self.analysis.resolutions.get(issue_id)
            saved_choice = saved.field_choices.get(conflict.conflict_id, {}) if saved and saved.action == RESOLUTION_FIELDS else {}
            mode = str(self.diff_boxes[row].currentData() or "")
            value = self.diff_inputs[row].text().strip() if mode == "CUSTOM" else ""
            if mode != str(saved_choice.get("mode") or "") or value != str(saved_choice.get("value") or ""):
                return 2
        return None

    def _check_unapplied(self) -> bool:
        step = self._unapplied_step()
        if step is None:
            return True
        self.tabs.setCurrentIndex(step)
        self.feedback.setText((
            "Có dòng đã sửa nhưng chưa kiểm tra lại.",
            "Có Bill hoặc số quyết toán đã chọn nhưng chưa ghép.",
            "Có cách xử lý số liệu chưa lưu.",
        )[step] + " Hãy dùng nút màu xanh trong bước này trước khi tiếp tục.")
        return False

    def _reanalyze_or_restore(
        self, old_payload: Dict[str, Any], old_resolutions: Dict[str, IssueResolution],
        old_drafts: tuple[dict, dict, dict] | None = None,
    ) -> bool:
        try:
            self.sandbox.reanalyze()
        except DailyImportError as exc:
            self.sandbox.payload = old_payload
            self.sandbox.resolutions = old_resolutions
            if old_drafts is not None:
                self.edit_drafts, self.choice_drafts, self.field_drafts = old_drafts
            self.refresh_ui(capture=False)
            notify.error(self, "Không kiểm tra lại được dữ liệu", str(exc))
            return False
        self._dirty = True
        return True

    def capture_drafts(self):
        for row, issue_id in enumerate(self.edit_ids):
            day = self.edit_table.item(row, 3)
            container = self.edit_table.item(row, 4)
            if day and container:
                self.edit_drafts[issue_id] = (day.text(), container.text())
        for issue_id, box in zip(self.choice_ids, self.choice_boxes):
            self.choice_drafts[issue_id] = str(box.currentData() or "")
        for (_issue_id, conflict), box, editor in zip(self.diff_rows, self.diff_boxes, self.diff_inputs):
            self.field_drafts[conflict.conflict_id] = (str(box.currentData() or ""), editor.text())

    def refresh_ui(self, *, capture: bool = True):
        if capture:
            self.capture_drafts()
        self._refreshing = True
        analysis = self.sandbox.analysis
        pending = analysis.unresolved_issues
        edit_count = sum(issue.kind in EDIT_KINDS for issue in pending)
        choice_count = sum(issue.kind in CHOOSE_KINDS for issue in pending)
        field_count = sum(issue.kind == ISSUE_FIELD_CONFLICT for issue in pending)
        self.summary.setText(
            f"Còn {len(pending)} việc cần quyết định  ·  "
            f"{edit_count} cần sửa thông tin  ·  {choice_count} cần chọn Bill/số quyết toán  ·  "
            f"{field_count} cần chọn số liệu  ·  {analysis.skipped_count} dòng để sau"
        )
        self.tabs.setTabText(0, f"1. Sửa thông tin ({edit_count})")
        self.tabs.setTabText(1, f"2. Chọn Bill / số quyết toán ({choice_count})")
        self.tabs.setTabText(2, f"3. Chọn số liệu ({field_count})")
        self._fill_edit_table()
        self._fill_choice_table()
        self._fill_diff_table()
        for table in (self.edit_table, self.choice_table, self.diff_table):
            if table.rowCount() and table.currentRow() < 0:
                table.selectRow(0)
        self._refreshing = False
        self.update_edit_context()
        self.update_choice_context()
        self.update_diff_context()
        self.update_edit_actions()
        self.update_choice_actions()
        self.update_field_actions()
        self._fit_choice_table()
        self._size_choice_columns()
        self._size_diff_columns()
        self.refresh_review()

    def _fill_edit_table(self):
        issues = self.current_issues(EDIT_KINDS)
        self.edit_ids = [issue.issue_id for issue in issues]
        self.edit_table.setRowCount(len(issues))
        for index, issue in enumerate(issues):
            row = self.source_row(issue)
            key, key_label = self.effective_day_key(issue)
            day, container = self.edit_drafts.get(issue.issue_id, (shown(row.get(key)), str(row.get("so_container") or "")))
            for col, value in enumerate((self.source_label(issue), user_issue_name(issue), key_label)):
                self.edit_table.setItem(index, col, read_only(value))
            self.edit_table.setItem(index, 3, QTableWidgetItem(day))
            self.edit_table.setItem(index, 4, QTableWidgetItem(container))
            if self._is_deferred(issue.issue_id):
                for col in (3, 4):
                    item = self.edit_table.item(index, col)
                    item.setFlags(item.flags() & ~Qt.ItemIsEditable)
            self.edit_table.setItem(index, 5, read_only("Xử lý sau" if self._is_deferred(issue.issue_id) else "Cần sửa"))
            self._paint_deferred(self.edit_table, index, self._is_deferred(issue.issue_id))

    def _fill_choice_table(self):
        issues = self.current_issues(CHOOSE_KINDS)
        self.choice_ids = [issue.issue_id for issue in issues]
        self.choice_boxes = []
        self.choice_table.setRowCount(len(issues))
        for index, issue in enumerate(issues):
            for col, value in enumerate((self.source_label(issue), user_issue_name(issue), shown(issue.match_date), issue.container)):
                self.choice_table.setItem(index, col, read_only(value))
            box = QComboBox()
            box.addItem("Chọn cho dòng này…", "")
            for candidate in issue.candidates:
                box.addItem(user_candidate_name(candidate), candidate.candidate_id)
            saved = self.sandbox.analysis.resolutions.get(issue.issue_id)
            wanted = self.choice_drafts.get(issue.issue_id) or (saved.candidate_id if saved and saved.action == RESOLUTION_SELECT else "")
            index_in_box = box.findData(wanted)
            if index_in_box >= 0:
                box.setCurrentIndex(index_in_box)
            box.currentIndexChanged.connect(lambda _value, row=index: self._choice_changed(row))
            self.choice_table.setCellWidget(index, 4, self._centered_table_control(box))
            self.choice_boxes.append(box)
            self.choice_table.setItem(index, 5, read_only("Xử lý sau" if saved and saved.action == RESOLUTION_SKIP else "Đã chọn" if saved else "Chưa chọn"))
            box.setEnabled(not self._is_deferred(issue.issue_id))
            if self._is_deferred(issue.issue_id):
                box.setStyleSheet("background: #FFF1B8; color: #785A00;")
            self._paint_deferred(self.choice_table, index, self._is_deferred(issue.issue_id))
    def _choice_changed(self, row: int):
        if self._refreshing:
            return
        issue_id = self.choice_ids[row]
        saved = self.sandbox.analysis.resolutions.get(issue_id)
        if saved and saved.action == RESOLUTION_SELECT and not self.choice_boxes[row].currentData():
            prior_index = self.choice_boxes[row].findData(saved.candidate_id)
            if prior_index > 0:
                blocked = self.choice_boxes[row].blockSignals(True)
                self.choice_boxes[row].setCurrentIndex(prior_index)
                self.choice_boxes[row].blockSignals(blocked)
        self.choice_table.selectRow(row)
        self.update_choice_context()
        self.update_choice_actions()

    def _fill_diff_table(self):
        pairs = [(issue, conflict) for issue in self.current_issues({ISSUE_FIELD_CONFLICT}) for conflict in issue.conflicts]
        self.diff_rows = [(issue.issue_id, conflict) for issue, conflict in pairs]
        self.diff_boxes = []
        self.diff_inputs = []
        self.diff_table.setRowCount(len(pairs))
        for row_index, (issue, conflict) in enumerate(pairs):
            values = (self.source_label(issue), self._issue_sqt(issue, conflict), conflict.field_name,
                      shown(conflict.current_value), shown(conflict.new_value))
            for col, value in enumerate(values):
                item = read_only(value)
                item.setToolTip(str(value) if value not in (None, "") else "")
                item.setTextAlignment(Qt.AlignLeft | Qt.AlignVCenter)
                if col == 4:
                    item.setForeground(QColor("#225BAD"))
                self.diff_table.setItem(row_index, col, item)
            box = QComboBox()
            box.setObjectName("conflictRule")
            box.addItem("Chưa chọn", "")
            box.addItem("Giữ số liệu đang có", "CURRENT")
            box.addItem("Dùng số liệu mới", "NEW")
            box.addItem("Nhập giá trị khác", "CUSTOM")
            editor = QLineEdit()
            editor.setObjectName("conflictCustom")
            saved = self.sandbox.analysis.resolutions.get(issue.issue_id)
            saved_choice = saved.field_choices.get(conflict.conflict_id, {}) if saved else {}
            mode, value = self.field_drafts.get(conflict.conflict_id, (str(saved_choice.get("mode") or ""), shown(saved_choice.get("value"))))
            box.setCurrentIndex(max(0, box.findData(mode)))
            editor.setText(value)
            editor.setEnabled(box.currentData() == "CUSTOM")
            box.currentIndexChanged.connect(lambda _index, target=editor, source=box: target.setEnabled(source.currentData() == "CUSTOM"))
            box.currentIndexChanged.connect(self.update_field_actions)
            editor.textChanged.connect(self.update_field_actions)
            self.diff_table.setCellWidget(row_index, 5, self._centered_table_control(box))
            self.diff_table.setCellWidget(row_index, 6, self._centered_table_control(editor))
            self.diff_boxes.append(box)
            self.diff_inputs.append(editor)
            self.diff_table.setItem(row_index, 7, read_only("Xử lý sau" if saved and saved.action == RESOLUTION_SKIP else "Đã chọn" if saved else "Chưa chọn"))
            box.setEnabled(not self._is_deferred(issue.issue_id))
            editor.setEnabled(not self._is_deferred(issue.issue_id) and box.currentData() == "CUSTOM")
            if self._is_deferred(issue.issue_id):
                box.setStyleSheet("background: #FFF1B8; color: #785A00;")
                editor.setStyleSheet("background: #FFF1B8;")
            self._paint_deferred(self.diff_table, row_index, self._is_deferred(issue.issue_id))
            if row_index % 2 and not self._is_deferred(issue.issue_id):
                for column in range(self.diff_table.columnCount()):
                    if item := self.diff_table.item(row_index, column):
                        item.setBackground(QColor("#FAFCFF"))
        old_field = self.field_filter.currentText()
        self.field_filter.blockSignals(True)
        self.field_filter.clear()
        self.field_filter.addItem("Tất cả")
        self.field_filter.addItems(sorted({conflict.field_name for _, conflict in pairs}))
        self.field_filter.setCurrentIndex(max(0, self.field_filter.findText(old_field)))
        self.field_filter.blockSignals(False)
        self.filter_fields()

    def filter_fields(self):
        wanted = self.field_filter.currentText()
        visible = 0
        for row, (_issue_id, conflict) in enumerate(self.diff_rows):
            show = wanted == "Tất cả" or wanted == conflict.field_name
            self.diff_table.setRowHidden(row, not show)
            visible += show
        self.scope.setText(f"Đang hiện {visible}/{len(self.diff_rows)} trường")
        self.update_rule_action()

    def fill_rule(self):
        mode = str(self.rule.currentData() or "")
        if not mode:
            QMessageBox.information(self, "Chưa chọn cách xử lý", "Hãy chọn cách xử lý trước khi điền vào bảng.")
            return
        count = 0
        for row, box in enumerate(self.diff_boxes):
            saved = self.analysis.resolutions.get(self.diff_rows[row][0])
            if not self.diff_table.isRowHidden(row) and not box.currentData() and not (saved and saved.action == RESOLUTION_SKIP):
                box.setCurrentIndex(box.findData(mode))
                count += 1
        self.update_field_actions()
        self.feedback.setText(f"Đã điền cách xử lý cho {count} trường đang hiện. Kiểm tra ngoại lệ rồi bấm Lưu dữ liệu.")

    def apply_edits(self):
        changes = []
        for index, issue_id in enumerate(self.edit_ids):
            issue = next(issue for issue in self.sandbox.analysis.issues if issue.issue_id == issue_id)
            row = self.source_row(issue)
            key, _label = self.effective_day_key(issue)
            day = self.edit_table.item(index, 3).text().strip()
            container = self.edit_table.item(index, 4).text().strip().upper()
            if (day, container) == (shown(row.get(key)), str(row.get("so_container") or "").upper()):
                continue
            if day and parse_date(day) is None:
                QMessageBox.warning(self, "Ngày chưa đúng", f"Dòng {index + 1}: hãy nhập ngày theo dạng 15/06/2026.")
                return
            changes.append((issue.json_indexes[0], key, day, container, issue_id))
        if not changes:
            self.feedback.setText("Chưa có ngày hoặc Container nào thay đổi.")
            return
        self.capture_drafts()
        old_payload = copy.deepcopy(self.sandbox.payload)
        old_resolutions = dict(self.sandbox.resolutions)
        old_drafts = (dict(self.edit_drafts), dict(self.choice_drafts), dict(self.field_drafts))
        for json_index, key, day, container, issue_id in changes:
            row = self.sandbox.payload["du_lieu_boc_tach"][json_index]
            row[key] = day
            row["so_container"] = container
            self.edit_drafts.pop(issue_id, None)
        old_ids = {issue.issue_id for issue in self.sandbox.analysis.unresolved_issues}
        if not self._reanalyze_or_restore(old_payload, old_resolutions, old_drafts):
            return
        new_ids = {issue.issue_id for issue in self.sandbox.analysis.unresolved_issues}
        if not self._persist_current():
            return
        self.refresh_ui(capture=False)
        self.feedback.setText(f"PM đã kiểm tra lại {len(changes)} dòng: {len(old_ids - new_ids)} vấn đề hết, {len(new_ids - old_ids)} vấn đề mới phát sinh.")
        self._advance_if_done(0)

    def apply_choices(self):
        selected = []
        for issue_id, box in zip(self.choice_ids, self.choice_boxes):
            candidate_id = str(box.currentData() or "")
            saved = self.sandbox.analysis.resolutions.get(issue_id)
            if candidate_id and (saved is None or saved.action != RESOLUTION_SELECT or saved.candidate_id != candidate_id):
                selected.append((issue_id, candidate_id))
        if not selected:
            self.feedback.setText("Chưa có lựa chọn mới nào cần áp dụng.")
            return
        self.capture_drafts()
        old_payload = self.sandbox.payload
        old_resolutions = dict(self.sandbox.resolutions)
        for issue_id, candidate_id in selected:
            self.sandbox.resolutions[issue_id] = IssueResolution(issue_id, RESOLUTION_SELECT, candidate_id=candidate_id)
        old_ids = {issue.issue_id for issue in self.sandbox.analysis.unresolved_issues}
        if not self._reanalyze_or_restore(old_payload, old_resolutions):
            return
        new_ids = {issue.issue_id for issue in self.sandbox.analysis.unresolved_issues}
        if not self._persist_current():
            return
        self.refresh_ui(capture=False)
        self.feedback.setText(f"PM đã ghép lại {len(selected)} chứng từ: {len(old_ids - new_ids)} vấn đề hết, {len(new_ids - old_ids)} vấn đề mới phát sinh. Bill còn lại giữ PENDING để xử lý sau.")
        self._advance_if_done(1)

    def _valid_custom(self, conflict, value: str) -> bool:
        reference = conflict.new_value if conflict.new_value not in (None, "") else conflict.current_value
        if isinstance(reference, (date, datetime)):
            return parse_date(value) is not None
        if isinstance(reference, (int, float)) and not isinstance(reference, bool):
            return parse_number(value) is not None
        return bool(value)

    def apply_fields(self):
        self.capture_drafts()
        to_apply = []
        for issue in self.current_issues({ISSUE_FIELD_CONFLICT}):
            choices = {}
            for conflict in issue.conflicts:
                mode, value = self.field_drafts.get(conflict.conflict_id, ("", ""))
                if mode == "CUSTOM" and not self._valid_custom(conflict, value):
                    QMessageBox.warning(self, "Giá trị chưa đúng", f"Hãy kiểm tra giá trị nhập cho trường ‘{conflict.field_name}’.")
                    return
                if mode:
                    choices[conflict.conflict_id] = {"mode": mode, "value": value}
            if len(choices) == len(issue.conflicts):
                saved = self.sandbox.analysis.resolutions.get(issue.issue_id)
                if saved is None or saved.action != RESOLUTION_FIELDS or saved.field_choices != choices:
                    to_apply.append((issue.issue_id, choices))
        if not to_apply:
            self.feedback.setText("Chọn đủ cách xử lý cho mọi trường của một chứng từ để áp dụng.")
            return
        old_payload = self.sandbox.payload
        old_resolutions = dict(self.sandbox.resolutions)
        for issue_id, choices in to_apply:
            self.sandbox.resolutions[issue_id] = IssueResolution(issue_id, RESOLUTION_FIELDS, field_choices=choices)
        if not self._reanalyze_or_restore(old_payload, old_resolutions):
            return
        if not self._persist_current():
            return
        self.refresh_ui(capture=False)
        self.feedback.setText(f"PM đã ghi nhận lựa chọn số liệu cho {len(to_apply)} chứng từ. File theo dõi chưa được ghi.")
        self._advance_if_done(2)

    def _group_selection(self, group: str) -> list[ImportIssue]:
        table, ids = {"edit": (self.edit_table, self.edit_ids),
                      "choice": (self.choice_table, self.choice_ids),
                      "diff": (self.diff_table, [issue_id for issue_id, _ in self.diff_rows])}[group]
        rows = self._selected_rows(table)
        if not rows and table.currentRow() >= 0:
            rows = [table.currentRow()]
        selected_ids = {ids[row] for row in rows if row < len(ids)}
        return [issue for issue in self.sandbox.analysis.issues if issue.issue_id in selected_ids]

    def defer_selected(self, group: str):
        selected = [issue for issue in self._group_selection(group) if not self._is_deferred(issue.issue_id)]
        if not selected:
            QMessageBox.information(self, "Chưa chọn dòng", "Hãy chọn chứng từ cần xử lý sau.")
            return
        self.capture_drafts()
        old_payload = self.sandbox.payload
        old_resolutions = dict(self.sandbox.resolutions)
        old_drafts = (dict(self.edit_drafts), dict(self.choice_drafts), dict(self.field_drafts))
        for issue in selected:
            deferred_id = ROW_DEFER_PREFIX + issue.issue_id
            metadata = {
                "kind": issue.kind,
                "doc_type": issue.doc_type,
                "container": issue.container,
                "reference": issue.reference,
                "match_date": issue.match_date,
                "json_indexes": issue.json_indexes,
                "conflicts": [
                    {"conflict_id": conflict.conflict_id, "target_row": conflict.target_row,
                     "field_name": conflict.field_name, "current_value": shown(conflict.current_value),
                     "new_value": shown(conflict.new_value), "target_sheet": conflict.target_sheet}
                    for conflict in issue.conflicts
                ],
            }
            self.sandbox.resolutions.pop(issue.issue_id, None)
            self.sandbox.resolutions[deferred_id] = IssueResolution(
                deferred_id, RESOLUTION_SKIP, field_choices=metadata,
            )
            for conflict in issue.conflicts:
                self.field_drafts.pop(conflict.conflict_id, None)
            self.choice_drafts.pop(issue.issue_id, None)
        if not self._reanalyze_or_restore(old_payload, old_resolutions, old_drafts):
            return
        if not self._persist_current():
            return
        self.refresh_ui(capture=False)
        deferred_rows = len({index for issue in selected for index in issue.json_indexes})
        self.feedback.setText(f"Đã đánh dấu {deferred_rows} dòng chứng từ xử lý sau. Chọn dòng vàng rồi bấm Hoàn tác để đưa trở lại.")

    def undo_defer(self, group: str = "diff"):
        selected = [issue for issue in self._group_selection(group) if self._is_deferred(issue.issue_id)]
        if not selected:
            return
        old_payload = self.sandbox.payload
        old_resolutions = dict(self.sandbox.resolutions)
        for issue in selected:
            self.sandbox.resolutions.pop(issue.issue_id, None)
        if not self._reanalyze_or_restore(old_payload, old_resolutions):
            return
        if not self._persist_current():
            return
        self.refresh_ui(capture=False)
        self.feedback.setText(f"Đã hoàn tác xử lý sau cho {len(selected)} chứng từ.")

    def on_tab_change(self, index: int):
        if index == 1:
            QTimer.singleShot(0, self._size_choice_columns)
        if index == 2:
            QTimer.singleShot(0, self._size_diff_columns)
        if index == 3:
            self.refresh_review()

    def refresh_review(self):
        analysis = self.sandbox.analysis
        waiting = len(analysis.unresolved_issues)
        unapplied = self._unapplied_step() is not None
        if waiting:
            review_state = f"Còn {waiting} việc cần quyết định trước khi xác nhận."
        elif unapplied:
            review_state = "Có thay đổi chưa áp dụng; hãy quay lại bước tương ứng."
        else:
            review_state = "Đã đủ quyết định để xác nhận."
        self.review_summary.setText(
            f"Dự kiến tạo {analysis.new_info_count} dòng quyết toán, cập nhật {analysis.updated_info_count} dòng quyết toán, "
            f"ghi {len(analysis.expense_changes)} dòng khoản chi; {analysis.skipped_count} dòng để sau. "
            + review_state
        )
        waiting_names = [user_issue_name(issue) for issue in analysis.unresolved_issues]
        skipped_names = []
        for index in sorted(set(analysis.skipped_json_indexes)):
            if 0 <= index < len(self.sandbox.payload["du_lieu_boc_tach"]):
                source = self.sandbox.payload["du_lieu_boc_tach"][index]
                identifier = source.get("so_bill") or source.get("so_hd") or source.get("file_nguon") or "Chứng từ"
                skipped_names.append(f"{source.get('loai_chung_tu') or 'Chứng từ'} {identifier}")
        details = []
        if waiting_names:
            details.append("<b>Cần quyết định:</b> " + html.escape(" · ".join(waiting_names)))
        if skipped_names:
            details.append("<b>Giữ PENDING để xử lý sau:</b> " + html.escape(" · ".join(skipped_names)))
        self.review_attention.setText("<br>".join(details) if details else "<b>Không có chứng từ nào cần xử lý sau.</b>")
        changes = []
        self._review_changes = []
        for change in analysis.info_changes:
            sheet = change.target_sheet or "Thông tin quyết toán"
            details = self._change_details(change, sheet, expense=False)
            position = f"{sheet}, dòng {change.target_row}" if change.target_row else f"{sheet}, dòng mới"
            changes.append(("Quyết toán", "Tạo dòng mới" if change.action == "CREATE" else "Cập nhật dòng",
                            change.sqt, change.values.get("Số Container"), position, self._change_summary(details)))
            self._review_changes.append((change, details, position))
        for change in analysis.expense_changes:
            details = self._change_details(change, EXPENSE_SHEET, expense=True)
            position = f"{EXPENSE_SHEET}, dòng {change.target_row}" if change.target_row else f"{EXPENSE_SHEET}, dòng mới"
            changes.append(("Khoản chi", "Tạo dòng mới" if change.action == "CREATE" else "Cập nhật dòng",
                            change.values.get("SQT PM"), change.values.get("Số Container"), position,
                            self._change_summary(details)))
            self._review_changes.append((change, details, position))
        self.review_table.setRowCount(len(changes))
        for row, values in enumerate(changes):
            for col, value in enumerate(values):
                self.review_table.setItem(row, col, read_only(value))
        if changes and self.review_table.currentRow() < 0:
            self.review_table.selectRow(0)
        self.update_review_detail()
        self.finish_button.setEnabled(waiting == 0 and not unapplied)

    @staticmethod
    def _change_summary(details: list[tuple[str, str, str]]) -> str:
        if not details:
            return "Không đổi ô dữ liệu"
        names = ", ".join(item[0] for item in details[:3])
        return f"{len(details)} trường: {names}" + ("…" if len(details) > 3 else "")

    def _change_details(self, change, sheet: str, *, expense: bool) -> list[tuple[str, str, str]]:
        current = self.sandbox.workbook_rows.get((sheet, change.target_row), {}) if change.target_row else {}
        conflict_by_name = {conflict.field_name: conflict for conflict in change.conflicts}
        internal = {"MD5", "Trạng thái nhập", "Ngày cập nhật", "Ngày nhập cuối"}
        details = []
        for name, new_value in change.values.items():
            if name in internal or new_value in (None, ""):
                continue
            old_value = current.get(name)
            if change.action == "UPDATE":
                conflict = conflict_by_name.get(name)
                if conflict and not self.analysis.conflict_decisions.get(conflict.conflict_id, False):
                    continue
                if not conflict and old_value not in (None, "") and not (expense and name == "Nơi đóng"):
                    continue
                if shown(old_value) == shown(new_value):
                    continue
            details.append((name, shown(old_value) or "Trống", shown(new_value)))
        return details

    def update_review_detail(self):
        row = self.review_table.currentRow()
        if not hasattr(self, "_review_changes") or not 0 <= row < len(self._review_changes):
            self.review_details.setHtml("Chọn một dòng ở bảng trên để xem các trường sẽ ghi.")
            return
        change, details, position = self._review_changes[row]
        sources = []
        md5s = set(change.document_md5s)
        for source in self.sandbox.payload.get("du_lieu_boc_tach") or []:
            if str(source.get("ma_md5_file") or "").strip().lower() in md5s:
                name = str(source.get("file_nguon") or "").strip()
                if name and name not in sources:
                    sources.append(name)
        title = "Tạo dòng mới" if change.action == "CREATE" else "Cập nhật dòng đang có"
        parts = [f"<b>{title}</b> · {html.escape(position)} · SQT {html.escape(shown(change.values.get('SQT PM')) or '—')}"]
        if sources:
            parts.append("<b>Chứng từ nguồn:</b> " + html.escape(" · ".join(sources)))
        if details:
            parts.append("<b>Các ô sẽ ghi:</b><br>" + "<br>".join(
                f"• <b>{html.escape(name)}</b>: {html.escape(old)} → {html.escape(new)}"
                for name, old, new in details
            ))
        else:
            parts.append("Không có ô dữ liệu nào thay đổi; chứng từ chỉ được ghi nhận vào thông tin theo dõi.")
        self.review_details.setHtml("<br><br>".join(parts))

    def finish(self):
        if not self._check_unapplied():
            return
        if self.analysis.unresolved_issues:
            self.feedback.setText("Hãy quyết định hoặc để sau mọi chứng từ trước khi xác nhận nhập.")
            return
        try:
            save_extract_payload(self.analysis.output_path, self.sandbox.payload)
            final = self.sandbox.service.analyze(
                self.analysis.output_path, self.analysis.daily_path,
                processed_file_id=self.analysis.processed_file_id,
                resolutions=self.sandbox.resolutions,
            )
        except DailyImportError as exc:
            notify.error(self, "Không chuẩn bị được lần nhập", str(exc))
            return
        if final.unresolved_issues:
            self.sandbox.analysis = final
            self.sandbox.resolutions = dict(final.resolutions)
            self.refresh_ui(capture=False)
            notify.warning(self, "Dữ liệu đã thay đổi", "Một số lựa chọn không còn hợp lệ. Hãy kiểm tra lại trước khi nhập.")
            return
        self.final_analysis = final
        self._dirty = False
        self.accept()

    def save_and_close(self):
        if not self._check_unapplied():
            return
        if not self._persist_current():
            return
        self.saved_only = True
        self.accept()

    def _persist_current(self) -> bool:
        """Save the current working decisions without writing the workbook."""
        try:
            save_extract_payload(self.analysis.output_path, self.sandbox.payload)
            saved = self.sandbox.service.analyze(
                self.analysis.output_path, self.analysis.daily_path,
                processed_file_id=self.analysis.processed_file_id,
                resolutions=self.sandbox.resolutions,
            )
            self.sandbox.service.save_resolution_draft(saved.output_path, saved.daily_path, saved.resolutions)
            self.sandbox.service.save_work_in_progress(saved, saved.resolutions)
            source_key, _ = self.sandbox.service._deferred_identity(saved)
            batch = self.sandbox.service.database.get_deferred_batch_by_key(source_key)
            if batch:
                saved.deferred_batch_id = batch["id"]
        except (DailyImportError, OSError) as exc:
            notify.error(self, "Không lưu được xử lý", str(exc))
            return False
        self.sandbox.analysis = saved
        self.sandbox.resolutions = dict(saved.resolutions)
        self._dirty = False
        self._has_saved_draft = True
        return True

    def reject(self):  # noqa: N802 - Qt API
        if (self._dirty or self._unapplied_step() is not None) and not notify.confirm(
            self, "Chưa lưu xử lý", "Đóng cửa sổ và bỏ các chỉnh sửa/lựa chọn chưa lưu?",
            confirm_text="Bỏ thay đổi", cancel_text="Tiếp tục xử lý",
            tone="warning", confirm_variant="danger",
        ):
            return
        if self._has_saved_draft:
            self.saved_only = True
            self.accept()
        else:
            super().reject()


