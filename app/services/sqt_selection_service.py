"""Read SQT choices from the daily workbook and write PAD selection JSON."""

from __future__ import annotations

import json
import math
import os
import re
import unicodedata
from collections import OrderedDict
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable, List

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .. import file_utils
from ..daily_import import (
    EXPENSE_SHEET,
    INFO_SHEET,
    is_month_info_sheet,
    parse_date,
    parse_number,
)

SQT_COLUMN = "SQT PM"
STATUS_COLUMN = "Trạng thái nhập"
SELECTION_OPERATION = "NHAP_THONG_TIN"
EXPENSE_SELECTION_OPERATION = "NHAP_KHOAN_CHI"


class SqtSelectionError(RuntimeError):
    """Business error safe to show directly to the user."""


class DailyFileNotFoundError(SqtSelectionError):
    pass


class DailyFileLockedError(SqtSelectionError):
    pass


class MissingSheetError(SqtSelectionError):
    pass


class MissingColumnError(SqtSelectionError):
    pass


class EmptySqtListError(SqtSelectionError):
    pass


class SelectionJsonWriteError(SqtSelectionError):
    pass


class MultipleSheetsSelectedError(SqtSelectionError):
    """SQT được chọn nằm ở nhiều sheet tháng khác nhau.

    Luồng PAD nhập thông tin mở đúng một sheet mỗi lần chạy, nên khi các SQT trải
    trên nhiều tháng, app chặn lại để người dùng chọn gọn trong một tháng.
    """

    def __init__(self, sheets: Iterable[str]):
        self.sheets = sorted({str(s) for s in sheets if str(s).strip()})
        super().__init__(
            "Các Số quyết toán đang thuộc nhiều tháng: " + ", ".join(self.sheets)
        )


@dataclass(frozen=True)
class SqtSourceRow:
    row_number: int
    container: str = ""
    close_date: str = ""
    imported: bool = False
    total_amount: float | None = None
    material_price: float | None = None


@dataclass(frozen=True)
class SqtSelectionItem:
    value: str
    row_count: int
    sheet_name: str = ""
    row_numbers: tuple[int, ...] = ()
    imported_row_count: int = 0
    source_rows: tuple[SqtSourceRow, ...] = ()
    month_sheet_name: str = ""

    @property
    def pending_row_count(self) -> int:
        return max(0, self.row_count - self.imported_row_count)

    @property
    def is_fully_imported(self) -> bool:
        return self.row_count > 0 and self.pending_row_count == 0

    @property
    def is_partially_imported(self) -> bool:
        return self.imported_row_count > 0 and self.pending_row_count > 0

    @property
    def display_text(self) -> str:
        if self.sheet_name:
            return f"{self.value} — {self.row_count} dòng — {self.sheet_name}"
        return f"{self.value} — {self.row_count} dòng"

    @property
    def status_kind(self) -> str:
        if self.is_fully_imported:
            return "done"
        return "partial" if self.is_partially_imported else "pending"

    @property
    def status_text(self) -> str:
        if self.is_fully_imported:
            return "Đã nhập"
        if self.is_partially_imported:
            return f"Đã nhập {self.imported_row_count}/{self.row_count}"
        return "Chưa nhập"

    def to_selection_payload(self) -> dict[str, Any]:
        payload: dict[str, Any] = {"sqt": self.value}
        if self.sheet_name:
            payload["sheet_name"] = self.sheet_name
        if self.row_numbers:
            payload["row_numbers"] = list(self.row_numbers)
        return payload


def normalize_sqt_value(value: Any) -> str:
    """Return the stable string sent to PAD for an Excel SQT cell value."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value).strip()
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            return ""
        if value.is_integer():
            return str(int(value))
        return format(value, "g").strip()

    text = str(value).strip()
    if not text:
        return ""
    if text.endswith(".0"):
        head = text[:-2]
        if head and head.lstrip("+-").isdigit():
            return str(int(head))
    return text


def is_imported_status(value: Any) -> bool:
    """Return whether an Excel import-status cell represents a completed row."""
    text = unicodedata.normalize("NFD", str(value or ""))
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    text = text.upper().replace("Đ", "D")
    key = re.sub(r"\s+", " ", text).strip()
    return key in {"DA NHAP", "HOAN THANH"}


def _header_map(ws) -> dict[str, int]:
    return {
        str(cell.value).strip(): cell.column
        for cell in ws[1]
        if cell.value not in (None, "")
    }


def _cell(row: tuple[Any, ...], column: int | None) -> Any:
    return row[column - 1] if column and column <= len(row) else None


def _text_cell(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def _money_cell(value: Any) -> float | None:
    number = parse_number(value)
    if number is None and isinstance(value, str):
        compact = re.sub(r"\s*(?:đ|vnđ)\s*$", "", value.strip(), flags=re.IGNORECASE)
        if re.fullmatch(r"-?\d{1,3}(?:\.\d{3})+", compact):
            number = float(compact.replace(".", ""))
    return number if number is not None and math.isfinite(number) else None


def _info_sheet_names(wb, requested_sheet: str) -> list[str]:
    if requested_sheet in wb.sheetnames:
        return [requested_sheet]
    if requested_sheet == INFO_SHEET:
        monthly = [name for name in wb.sheetnames if is_month_info_sheet(name)]
        if monthly:
            return monthly
    raise MissingSheetError(f"Không tìm thấy sheet “{requested_sheet}”.")


def read_sqt_items(daily_file: str, sheet_name: str = INFO_SHEET) -> List[SqtSelectionItem]:
    """Read unique SQT values in first-seen order and count source rows."""
    if not daily_file or not os.path.isfile(daily_file):
        raise DailyFileNotFoundError("Không tìm thấy file Excel hằng ngày.")

    try:
        if file_utils.is_file_locked(daily_file):
            raise DailyFileLockedError(
                "File Excel hằng ngày đang bị khóa. Vui lòng đóng Excel và thử lại."
            )
    except FileNotFoundError as exc:
        raise DailyFileNotFoundError("Không tìm thấy file Excel hằng ngày.") from exc

    wb = None
    try:
        wb = load_workbook(daily_file, read_only=True, data_only=True)
        sheet_names = _info_sheet_names(wb, sheet_name)

        expense_months: dict[str, set[str]] = {}
        if sheet_name == EXPENSE_SHEET:
            for month_sheet in (name for name in wb.sheetnames if is_month_info_sheet(name)):
                month_ws = wb[month_sheet]
                month_sqt_col = _header_map(month_ws).get(SQT_COLUMN)
                if month_sqt_col is None:
                    continue
                for (raw_sqt,) in month_ws.iter_rows(
                    min_row=2, min_col=month_sqt_col, max_col=month_sqt_col, values_only=True
                ):
                    month_sqt = normalize_sqt_value(raw_sqt)
                    if month_sqt:
                        expense_months.setdefault(month_sqt, set()).add(month_sheet)

        counts: "OrderedDict[tuple[str, str], dict[str, Any]]" = OrderedDict()
        for current_sheet in sheet_names:
            ws = wb[current_sheet]
            headers = _header_map(ws)
            if SQT_COLUMN not in headers:
                raise MissingColumnError(
                    f"Không tìm thấy cột “{SQT_COLUMN}” trong sheet “{current_sheet}”."
                )

            sqt_col = headers[SQT_COLUMN]
            status_col = headers.get(STATUS_COLUMN)
            container_col = headers.get("Số Container")
            date_col = (
                headers.get("Ngày tháng") if current_sheet == EXPENSE_SHEET
                else headers.get("Ngày Đóng", headers.get("Ngày đóng"))
            )
            total_col = headers.get("Tổng tiền") if current_sheet == EXPENSE_SHEET else None
            material_col = headers.get("Giá vật liệu") if current_sheet == EXPENSE_SHEET else None
            max_col = max(
                col for col in (sqt_col, status_col, container_col, date_col, total_col, material_col)
                if col is not None
            )
            for row_number, row in enumerate(
                ws.iter_rows(
                    min_row=2,
                    min_col=1,
                    max_col=max_col,
                    values_only=True,
                ),
                start=2,
            ):
                sqt = normalize_sqt_value(_cell(row, sqt_col))
                if not sqt:
                    continue
                key = (current_sheet if len(sheet_names) > 1 else "", sqt)
                entry = counts.setdefault(
                    key,
                    {"row_numbers": [], "imported_row_count": 0, "source_rows": []},
                )
                entry["row_numbers"].append(row_number)
                imported = is_imported_status(_cell(row, status_col))
                if imported:
                    entry["imported_row_count"] += 1
                raw_date = _cell(row, date_col)
                entry["source_rows"].append(
                    SqtSourceRow(
                        row_number=row_number,
                        container=_text_cell(_cell(row, container_col)),
                        close_date=parse_date(raw_date) or _text_cell(raw_date),
                        imported=imported,
                        total_amount=_money_cell(_cell(row, total_col)),
                        material_price=_money_cell(_cell(row, material_col)),
                    )
                )

        if not counts:
            raise EmptySqtListError("Sheet không có SQT hợp lệ.")
        return [
            SqtSelectionItem(
                value=sqt,
                row_count=len(entry["row_numbers"]),
                sheet_name=sheet,
                row_numbers=tuple(entry["row_numbers"]),
                imported_row_count=entry["imported_row_count"],
                source_rows=tuple(entry["source_rows"]),
                month_sheet_name=(
                    next(iter(expense_months[sqt]))
                    if len(expense_months.get(sqt, ())) == 1 else ""
                ),
            )
            for (sheet, sqt), entry in counts.items()
        ]
    except SqtSelectionError:
        raise
    except PermissionError as exc:
        raise DailyFileLockedError(
            "File Excel hằng ngày đang bị khóa. Vui lòng đóng Excel và thử lại."
        ) from exc
    except (InvalidFileException, OSError, ValueError) as exc:
        raise SqtSelectionError(
            f"Không thể đọc file {os.path.basename(daily_file)}. "
            "Vui lòng đóng Excel và thử lại."
        ) from exc
    finally:
        if wb is not None:
            wb.close()


def resolve_selection_sheet(
    default_sheet: str,
    selected_items: Iterable[SqtSelectionItem] | None,
) -> str:
    """Tên sheet vật lý mà PAD sẽ mở cho lần nhập thông tin này.

    File theo tháng đã gắn tên sheet thật vào từng item (vd "Tháng 7"); file một
    sheet cũ để trống nên lấy ``default_sheet`` (``Thong_Tin_Quyet_Toan``). Nếu các
    SQT được chọn nằm ở nhiều tháng thì báo ``MultipleSheetsSelectedError`` vì
    luồng PAD chỉ nhập đúng một sheet mỗi lần chạy.
    """
    sheets = sorted(
        {item.sheet_name for item in (selected_items or []) if item.sheet_name}
    )
    if len(sheets) > 1:
        raise MultipleSheetsSelectedError(sheets)
    return sheets[0] if sheets else default_sheet


def read_selection_snapshot(
    daily_file: str,
    sheet_name: str,
    selected_items: Iterable[SqtSelectionItem],
) -> list[dict[str, Any]]:
    """Freeze the Excel rows selected for a PAD launch before PAD can edit them."""
    if file_utils.is_file_locked(daily_file):
        raise DailyFileLockedError(
            "File Excel đang được sử dụng hoặc có file khóa cần kiểm tra."
        )
    items = list(selected_items)
    selected = {item.value for item in items}
    selected_rows = {
        row_number for item in items for row_number in item.row_numbers
        if not item.sheet_name or item.sheet_name == sheet_name
    } if sheet_name != EXPENSE_SHEET else set()
    try:
        workbook = load_workbook(daily_file, read_only=True, data_only=True)
        try:
            if sheet_name not in workbook.sheetnames:
                raise MissingSheetError(f"Không tìm thấy sheet “{sheet_name}”.")
            sheet = workbook[sheet_name]
            headers = _header_map(sheet)
            sqt_col = headers.get(SQT_COLUMN)
            if sqt_col is None:
                raise MissingColumnError(f"Không tìm thấy cột “{SQT_COLUMN}”.")
            snapshot = []
            for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
                if selected_rows and row_number not in selected_rows:
                    continue
                sqt = normalize_sqt_value(_cell(row, sqt_col))
                if sqt not in selected:
                    continue
                values = {
                    name: _cell(row, column)
                    for name, column in headers.items()
                    if name not in {"MD5", "Ngày cập nhật"}
                }
                snapshot.append({
                    "sheet_name": sheet_name,
                    "row_number": row_number,
                    "sqt": sqt,
                    "values": values,
                })
            found = {item["sqt"] for item in snapshot}
            found_rows = {item["row_number"] for item in snapshot}
            if found != selected or (selected_rows and found_rows != selected_rows):
                raise SqtSelectionError(
                    "Một số SQT đã chọn không còn ở đúng dòng Excel. Hãy mở lại danh sách SQT."
                )
            return snapshot
        finally:
            workbook.close()
    except SqtSelectionError:
        raise
    except (InvalidFileException, OSError, ValueError) as exc:
        raise SqtSelectionError(f"Không đọc được dữ liệu đã chọn trong Excel: {exc}") from exc


def write_selection_json(
    target_path: str | os.PathLike[str],
    daily_file: str,
    selected_sqt: Iterable[str],
    sheet_name: str = INFO_SHEET,
    *,
    operation: str = SELECTION_OPERATION,
    selected_items: Iterable[SqtSelectionItem] | None = None,
) -> str:
    """Write the selected SQT list atomically for PAD to consume."""
    selected = [str(value) for value in selected_sqt if str(value).strip()]
    if not selected:
        raise SelectionJsonWriteError("Vui lòng chọn ít nhất một Số quyết toán.")

    target = Path(target_path)
    try:
        target.parent.mkdir(parents=True, exist_ok=True)
    except OSError as exc:
        raise SelectionJsonWriteError("Không tạo được thư mục runtime.") from exc

    payload = {
        "operation": str(operation or SELECTION_OPERATION).strip() or SELECTION_OPERATION,
        "daily_file": str(Path(daily_file).resolve()),
        "sheet_name": sheet_name,
        "selected_sqt": selected,
    }
    item_payloads = [
        item.to_selection_payload()
        for item in (selected_items or [])
        if item.value in selected
    ]
    if item_payloads and operation == SELECTION_OPERATION:
        payload["mode"] = "monthly_info"
        payload["selected_items"] = item_payloads

    temp_path = target.with_name(target.name + ".tmp")
    try:
        with open(temp_path, "w", encoding="utf-8") as f:
            json.dump(payload, f, ensure_ascii=False, indent=2)
            f.write("\n")
        os.replace(temp_path, target)
    except OSError as exc:
        try:
            if temp_path.exists():
                temp_path.unlink()
        except OSError:
            pass
        raise SelectionJsonWriteError("Không ghi được JSON lựa chọn RPA.") from exc

    return str(target.resolve())
