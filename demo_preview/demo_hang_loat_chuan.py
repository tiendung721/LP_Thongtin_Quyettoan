"""Demo xử lý xung đột bằng đúng bộ phân tích của PM, trên dữ liệu tạm.

Chạy: python demo_hang_loat_chuan.py
Toàn bộ JSON, Excel và database mẫu nằm trong thư mục tạm và bị dọn khi đóng.
"""

from __future__ import annotations

import copy
import html
import json
import logging
import sys
import tempfile
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from PySide6.QtCore import Qt
from PySide6.QtGui import QColor, QFontDatabase, QKeySequence
from PySide6.QtWidgets import (
    QApplication, QComboBox, QDialog, QHBoxLayout, QHeaderView, QLabel,
    QLineEdit, QMessageBox, QPushButton, QTableWidget, QTableWidgetItem,
    QTabWidget, QVBoxLayout, QWidget,
)

from app.daily_import import (
    DOC_EXPENSE, EXPENSE_HEADERS, EXPENSE_SHEET, INFO_HEADERS,
    ISSUE_FIELD_CONFLICT, ISSUE_MISSING_CLOSE_DATE, ISSUE_MULTIPLE_BILLS,
    ISSUE_MULTIPLE_TARGETS, ISSUE_NO_TARGET, RESOLUTION_FIELDS,
    RESOLUTION_SELECT, RESOLUTION_SKIP, DailyImportService, ImportIssue,
    IssueResolution, parse_date, parse_number,
)
from app.database import Database


EDIT_KINDS = {ISSUE_MISSING_CLOSE_DATE, ISSUE_NO_TARGET}
CHOOSE_KINDS = {ISSUE_MULTIPLE_BILLS, ISSUE_MULTIPLE_TARGETS}
STYLE = """
QWidget { font-family: 'Segoe UI', Arial; font-size: 10pt; color: #1D304A; }
QDialog { background: #F4F7FB; }
QLabel#title { font-size: 18pt; font-weight: 700; color: #102541; }
QLabel#heading { font-size: 12pt; font-weight: 700; color: #102541; }
QLabel#hint { color: #5E718B; }
QLabel#summary { background: white; border: 1px solid #DCE5F0; border-radius: 7px; padding: 8px 10px; color: #315170; }
QLabel#notice { background: #EAF3FF; color: #245A9D; padding: 9px; border-radius: 7px; }
QLabel#detail { background: #F6F9FD; border: 1px solid #DCE5F0; border-radius: 7px; padding: 10px; }
QPushButton { background: white; border: 1px solid #C6D3E2; border-radius: 7px; padding: 7px 12px; min-height: 23px; }
QPushButton:hover { background: #EDF3FB; }
QPushButton:disabled { color: #98A7B8; background: #F2F4F7; }
QPushButton[primary='true'] { color: white; background: #285DB7; border-color: #285DB7; font-weight: 700; }
QPushButton[quiet='true'] { color: #526983; background: #F8FAFD; border-color: #D6E0EC; }
QLineEdit, QComboBox { background: white; border: 1px solid #C7D3E2; border-radius: 6px; padding: 5px; }
QTableWidget { background: white; alternate-background-color: #FAFCFF; gridline-color: #E4EAF2; selection-background-color: #DDEBFF; selection-color: #172C4A; }
QHeaderView::section { background: #EDF2F8; padding: 8px; border: 0; border-right: 1px solid #DCE5F0; }
QTabWidget::pane { border: 1px solid #DCE5F0; background: white; }
QTabBar::tab { background: #E9EFF7; padding: 10px 16px; margin-right: 3px; }
QTabBar::tab:selected { background: white; color: #285DB7; font-weight: 700; }
"""


def label(value: str, role: str = "") -> QLabel:
    widget = QLabel(value)
    widget.setWordWrap(True)
    if role:
        widget.setObjectName(role)
    return widget


def button(value: str, action, primary: bool = False) -> QPushButton:
    widget = QPushButton(value)
    if primary:
        widget.setProperty("primary", "true")
    widget.clicked.connect(action)
    return widget


def read_only(value: Any) -> QTableWidgetItem:
    result = QTableWidgetItem(str(value) if value not in (None, "") else "—")
    result.setFlags(result.flags() & ~Qt.ItemIsEditable)
    return result


def shown(value: Any) -> str:
    if isinstance(value, (datetime, date)):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        number = float(value)
        if number.is_integer():
            return f"{int(number):,}".replace(",", ".")
        return f"{number:,.2f}".replace(",", "_").replace(".", ",").replace("_", ".").rstrip("0").rstrip(",")
    if isinstance(value, str):
        try:
            return datetime.strptime(value, "%Y-%m-%d").strftime("%d/%m/%Y")
        except ValueError:
            pass
    return str(value) if value not in (None, "") else ""


class PasteTable(QTableWidget):
    def keyPressEvent(self, event):  # noqa: N802 - Qt API
        if event.matches(QKeySequence.Paste):
            values = QApplication.clipboard().text().replace("\r\n", "\n").strip("\n")
            start_row, start_col = self.currentRow(), self.currentColumn()
            if start_row >= 0 and start_col in (3, 4):
                for row_offset, line in enumerate(values.split("\n")):
                    for col_offset, value in enumerate(line.split("\t")):
                        row, col = start_row + row_offset, start_col + col_offset
                        if row < self.rowCount() and col in (3, 4):
                            item = QTableWidgetItem(value.strip())
                            self.setItem(row, col, item)
                return
        super().keyPressEvent(event)


class DemoSandbox:
    """Cùng DailyImportService của PM; chỉ dùng file tạm của demo."""

    def __init__(self):
        self.temp = tempfile.TemporaryDirectory(prefix="pm_conflict_demo_")
        root = Path(self.temp.name)
        self.workbook_path = root / "theo_doi_mau.xlsx"
        self.json_path = root / "chung_tu_mau.json"
        self.database = Database(str(root / "trang_thai_mau.db"))
        self.database.init_db()
        self.service = DailyImportService(self.database, logging.getLogger("pm-conflict-demo"))
        self._create_workbook()
        self.workbook_rows = self._read_workbook_rows()
        self.original_payload = self._create_payload()
        self.payload = copy.deepcopy(self.original_payload)
        self.resolutions: dict[str, IssueResolution] = {}
        self.analysis = self.reanalyze()

    def _create_workbook(self):
        book = Workbook()
        info = book.active
        info.title = "Tháng 6"
        info.append(INFO_HEADERS)

        def info_row(sqt: int, day: int, container: str, **values):
            data = {"SQT PM": sqt, "Ngày Đóng": datetime(2026, 6, day), "Số Container": container}
            data.update(values)
            info.append([data.get(header) for header in INFO_HEADERS])

        info_row(210, 10, "AABB1234567", **{"Số tấn": 28, "Nơi đóng": "Kho Đông", "Loại hàng": "Gạo", "Người nhận": "Công ty An Phú"})
        info_row(211, 10, "AABB1234567", **{"Số tấn": 28, "Nơi đóng": "Kho Tây", "Loại hàng": "Gạo", "Người nhận": "Công ty Minh Long"})
        info_row(220, 11, "CCDD1234567", **{"Số tấn": 28, "Nơi đóng": "Kho A"})
        info_row(230, 12, "EEFF1234567", **{"Nơi đóng": "Kho C", "Loại hàng": "Gạo", "Người nhận": "Công ty Nam Việt"})
        info_row(231, 12, "EEFF1234567", **{"Nơi đóng": "Kho D", "Loại hàng": "Cám", "Người nhận": "Công ty Bình Minh"})
        info_row(240, 13, "GGHH1234567", **{"Nơi đóng": "Kho B"})
        info_row(250, 15, "IIJJ1234567")

        expense = book.create_sheet(EXPENSE_SHEET)
        expense.append(EXPENSE_HEADERS)
        existing = {
            "SQT PM": 240, "Ngày tháng": datetime(2026, 6, 13),
            "Số Container": "GGHH1234567", "Nơi đóng": "Kho B",
            "Số HĐ": "HD-99", "Đơn giá": 100_000,
        }
        expense.append([existing.get(header) for header in EXPENSE_HEADERS])
        book.save(self.workbook_path)
        book.close()

    def _create_payload(self) -> dict[str, Any]:
        rows = [
            # A1: một Phiếu cân có hai Bill cùng Container, khác MD5.
            dict(file_nguon="phieu_can_01.jpg", ma_md5_file="pc-a1", loai_chung_tu="Phiếu cân", trang_thai="OK",
                 ngay_dong="17/06/2026", so_container="MMNN1234567", so_chi_seal="SEAL-01", so_tan=29),
            dict(file_nguon="bill_01.pdf", ma_md5_file="bill-a1-1", loai_chung_tu="Bill", trang_thai="OK",
                 so_container="MMNN1234567", so_bill="BL-001", ngay_chay="18/06/2026", ten_tau="Tàu An Phú", vt_bien="Hãng Biển Xanh"),
            dict(file_nguon="bill_02.pdf", ma_md5_file="bill-a1-2", loai_chung_tu="Bill", trang_thai="OK",
                 so_container="MMNN1234567", so_bill="BL-002", ngay_chay="19/06/2026", ten_tau="Tàu Minh Long", vt_bien="Hãng Biển Nam"),
            # A2: Phiếu cân khớp hai dòng quyết toán cùng ngày và Container.
            dict(file_nguon="phieu_can_02.jpg", ma_md5_file="pc-a2", loai_chung_tu="Phiếu cân", trang_thai="OK",
                 ngay_dong="10/06/2026", so_container="AABB1234567", so_tan=28, noi_dong="Kho Đông"),
            # A3: thiếu Ngày đóng trong workbook chia sheet theo tháng.
            dict(file_nguon="phieu_can_03.jpg", ma_md5_file="pc-a3", loai_chung_tu="Phiếu cân", trang_thai="OK",
                 so_container="KKLL1234567", so_tan=27),
            # A4: dữ liệu Phiếu cân khác Excel ở trường Số tấn.
            dict(file_nguon="phieu_can_04.jpg", ma_md5_file="pc-a4", loai_chung_tu="Phiếu cân", trang_thai="OK",
                 ngay_dong="11/06/2026", so_container="CCDD1234567", so_tan=30, noi_dong="Kho A"),
            # B1: Khoản chi dùng Ngày chạy 14/06, trong Excel có SQT ngày 15/06.
            dict(file_nguon="chi_phi_01.jpg", ma_md5_file="cost-b1", loai_chung_tu="Khoản chi", trang_thai="OK",
                 ngay_dong="15/06/2026", ngay_chay="14/06/2026", so_container="IIJJ1234567", so_hd="HD-88", don_gia=110_000),
            # B2: một Khoản chi khớp hai SQT khác nhau.
            dict(file_nguon="chi_phi_02.jpg", ma_md5_file="cost-b2", loai_chung_tu="Khoản chi", trang_thai="OK",
                 ngay_chay="12/06/2026", so_container="EEFF1234567", so_hd="HD-89", don_gia=115_000, noi_dong="Kho C"),
            # B3: Khoản chi đã có cùng ngày + Container, nhưng Đơn giá khác.
            dict(file_nguon="chi_phi_03.jpg", ma_md5_file="cost-b3", loai_chung_tu="Khoản chi", trang_thai="OK",
                 ngay_chay="13/06/2026", so_container="GGHH1234567", so_hd="HD-99", don_gia=120_000),
        ]
        payload = {"metadata": {"phien_ban_schema": "1.0", "tong_so_dong_boc_tach": len(rows)},
                   "du_lieu_boc_tach": rows, "canh_bao": [], "raw_data": []}
        self.json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        return payload

    def _read_workbook_rows(self) -> dict[tuple[str, int], dict[str, Any]]:
        result = {}
        book = load_workbook(self.workbook_path, read_only=True, data_only=True)
        try:
            for sheet in book.worksheets:
                rows = sheet.iter_rows(values_only=True)
                headers = next(rows, ())
                for row_number, values in enumerate(rows, start=2):
                    if any(value is not None for value in values):
                        result[(sheet.title, row_number)] = dict(zip(headers, values))
        finally:
            book.close()
        return result

    def reanalyze(self):
        self.analysis = self.service.analyze(
            str(self.json_path), str(self.workbook_path),
            payload_override=self.payload, resolutions=self.resolutions,
        )
        return self.analysis

    def reset(self):
        self.payload = copy.deepcopy(self.original_payload)
        self.resolutions = {}
        self.reanalyze()

    def close(self):
        self.temp.cleanup()


def user_issue_name(issue: ImportIssue) -> str:
    if issue.kind == ISSUE_MULTIPLE_BILLS:
        return "A1 · Chọn Bill"
    if issue.kind == ISSUE_MULTIPLE_TARGETS:
        return "B2 · Chọn số quyết toán" if issue.doc_type == DOC_EXPENSE else "A2 · Chọn số quyết toán"
    if issue.kind == ISSUE_MISSING_CLOSE_DATE:
        return "A3 · Bổ sung ngày đóng"
    if issue.kind == ISSUE_NO_TARGET:
        return "B1 · Chưa tìm thấy số quyết toán"
    if issue.kind == ISSUE_FIELD_CONFLICT:
        return "B3 · Số liệu khác file theo dõi" if issue.doc_type == DOC_EXPENSE else "A4 · Số liệu khác file theo dõi"
    return issue.title


def user_candidate_name(candidate) -> str:
    if candidate.kind == "BILL":
        bill = candidate.values.get("Số B/L") or "chưa có số"
        vessel = candidate.values.get("Tên tàu") or "chưa có tên tàu"
        return f"Bill {bill} · {vessel} · dòng {candidate.json_index + 1}"
    return f"Số quyết toán {candidate.sqt} · {candidate.sheet_name}, dòng {candidate.row_number}"


class StandardBulkDemo(QDialog):
    def __init__(self):
        super().__init__()
        self.sandbox = DemoSandbox()
        self.setWindowTitle("Demo mới · Xử lý trước khi nhập file theo dõi")
        self.resize(1390, 860)
        self.setMinimumSize(1100, 710)
        self.edit_drafts: dict[str, tuple[str, str]] = {}
        self.choice_drafts: dict[str, str] = {}
        self.field_drafts: dict[str, tuple[str, str]] = {}
        self.edit_ids: list[str] = []
        self.choice_ids: list[str] = []
        self.diff_rows: list[tuple[str, Any]] = []
        self.choice_boxes: list[QComboBox] = []
        self.diff_boxes: list[QComboBox] = []
        self.diff_inputs: list[QLineEdit] = []
        self._refreshing = False
        self._build()
        self.refresh_ui(capture=False)

    def _build(self):
        outer = QVBoxLayout(self)
        outer.addWidget(label("Xử lý dữ liệu trước khi nhập", "title"))
        outer.addWidget(label("BẢN XEM THỬ · Dữ liệu mẫu · Chưa ghi vào file theo dõi", "notice"))
        self.summary = label("", "summary")
        outer.addWidget(self.summary)
        self.tabs = QTabWidget()
        self.tabs.addTab(self._edit_page(), "1. Sửa thông tin")
        self.tabs.addTab(self._choice_page(), "2. Chọn Bill / số quyết toán")
        self.tabs.addTab(self._diff_page(), "3. Chọn số liệu")
        self.tabs.addTab(self._review_page(), "4. Xem trước")
        self.tabs.currentChanged.connect(self.on_tab_change)
        outer.addWidget(self.tabs, 1)
        footer = QHBoxLayout()
        self.feedback = label("Chọn một nhóm để bắt đầu. Các mục có thể chuyển nhóm sau khi PM kiểm tra lại.", "hint")
        footer.addWidget(self.feedback, 1)
        self.undo_button = button("Hoàn tác để sau", self.undo_defer)
        self.undo_button.setVisible(False)
        footer.addWidget(self.undo_button)
        reset_button = button("Làm lại demo", self.reset)
        reset_button.setProperty("quiet", "true")
        footer.addWidget(reset_button)
        footer.addWidget(button("Xem trước kết quả", lambda: self.tabs.setCurrentIndex(3)))
        outer.addLayout(footer)
        self.last_deferred: tuple[str, IssueResolution | None] | None = None

    def _page(self, heading: str, hint: str):
        page = QWidget()
        layout = QVBoxLayout(page)
        layout.setContentsMargins(16, 16, 16, 14)
        layout.addWidget(label(heading, "heading"))
        layout.addWidget(label(hint, "hint"))
        return page, layout

    def _edit_page(self):
        page, layout = self._page(
            "Sửa thông tin để PM tìm đúng dòng",
            "Sửa ngày và Container ngay trong bảng. Với Khoản chi, cột ‘Ngày PM đang dùng’ cho biết phải sửa Ngày chạy hay Ngày đóng. Có thể dán nhiều ô cùng lúc.",
        )
        self.edit_table = PasteTable(0, 6)
        self.edit_table.setHorizontalHeaderLabels(["Chứng từ", "Vấn đề", "Ngày PM đang dùng", "Ngày", "Container", "Kết quả"])
        self.edit_table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.edit_table.verticalHeader().setVisible(False)
        self.edit_table.verticalHeader().setDefaultSectionSize(36)
        self.edit_table.setAlternatingRowColors(True)
        layout.addWidget(self.edit_table, 1)
        self.edit_detail, self.edit_preview = self._add_detail_panel(layout, "Dữ liệu liên quan trong file theo dõi")
        self.edit_table.itemSelectionChanged.connect(self.update_edit_context)
        self.edit_table.itemChanged.connect(self.update_edit_actions)
        actions = QHBoxLayout()
        sample = button("Demo: điền ví dụ", self.fill_edit_example)
        sample.setProperty("quiet", "true")
        actions.addWidget(sample)
        self.edit_defer = button("Để chứng từ này xử lý sau", lambda: self.defer_selected("edit"))
        self.edit_defer.setProperty("quiet", "true")
        actions.addWidget(self.edit_defer)
        actions.addStretch()
        self.edit_apply = button("Kiểm tra lại dòng đã sửa", self.apply_edits, True)
        actions.addWidget(self.edit_apply)
        layout.addLayout(actions)
        return page

    def _choice_page(self):
        page, layout = self._page(
            "Chọn Bill hoặc dòng quyết toán cho từng chứng từ",
            "Mỗi dòng có các lựa chọn riêng. Sau khi chọn, PM sẽ ghép lại và kiểm tra tiếp xem số liệu có khác file theo dõi không.",
        )
        self.choice_table = QTableWidget(0, 6)
        self.choice_table.setHorizontalHeaderLabels(["Chứng từ", "Vấn đề", "Ngày", "Container", "Chọn kết quả phù hợp", "Kết quả"])
        self.choice_table.horizontalHeader().setSectionResizeMode(QHeaderView.Interactive)
        self.choice_table.horizontalHeader().setStretchLastSection(True)
        for col, width in enumerate((205, 220, 105, 140, 440)):
            self.choice_table.setColumnWidth(col, width)
        self.choice_table.verticalHeader().setVisible(False)
        self.choice_table.verticalHeader().setDefaultSectionSize(36)
        self.choice_table.setAlternatingRowColors(True)
        self.choice_table.setSelectionBehavior(QTableWidget.SelectRows)
        self.choice_table.setEditTriggers(QTableWidget.NoEditTriggers)
        layout.addWidget(self.choice_table, 1)
        self.choice_detail, self.choice_preview = self._add_detail_panel(layout, "So sánh các phương án — bấm vào dòng để chọn")
        self.choice_table.itemSelectionChanged.connect(self.update_choice_context)
        self.choice_preview.cellClicked.connect(self.choose_from_preview)
        self.choice_hint = label("Bill không được chọn sẽ bỏ qua trong lần nhập này và giữ PENDING để xử lý sau.", "notice")
        layout.addWidget(self.choice_hint)
        actions = QHBoxLayout()
        sample = button("Demo: chọn ví dụ", self.fill_choice_example)
        sample.setProperty("quiet", "true")
        actions.addWidget(sample)
        self.choice_defer = button("Để chứng từ này xử lý sau", lambda: self.defer_selected("choice"))
        self.choice_defer.setProperty("quiet", "true")
        actions.addWidget(self.choice_defer)
        actions.addStretch()
        self.choice_apply = button("Ghép các lựa chọn", self.apply_choices, True)
        actions.addWidget(self.choice_apply)
        layout.addLayout(actions)
        return page

    def _diff_page(self):
        page, layout = self._page(
            "Chọn số liệu sẽ dùng trong file theo dõi",
            "Xem số liệu đang có và số liệu mới cạnh nhau. Có thể đặt một cách xử lý cho các trường đang hiện, rồi sửa ngoại lệ ở từng dòng.",
        )
        toolbar = QHBoxLayout()
        toolbar.addWidget(label("Xem trường:"))
        self.field_filter = QComboBox()
        self.field_filter.addItem("Tất cả")
        self.field_filter.currentIndexChanged.connect(self.filter_fields)
        toolbar.addWidget(self.field_filter)
        toolbar.addWidget(label("Cách xử lý cho các dòng đang hiện:"))
        self.rule = QComboBox()
        self.rule.addItem("Chọn cách xử lý", "")
        self.rule.addItem("Giữ số liệu đang có", "CURRENT")
        self.rule.addItem("Dùng số liệu mới", "NEW")
        self.rule.currentIndexChanged.connect(self.update_rule_action)
        toolbar.addWidget(self.rule)
        self.rule_fill = button("Điền cho các dòng đang hiện", self.fill_rule)
        toolbar.addWidget(self.rule_fill)
        toolbar.addStretch()
        layout.addLayout(toolbar)
        self.diff_table = QTableWidget(0, 7)
        self.diff_table.setHorizontalHeaderLabels(["Chứng từ", "Trường", "File theo dõi", "Chứng từ mới", "Cách xử lý", "Giá trị khác (nếu nhập)", "Kết quả"])
        self.diff_table.horizontalHeader().setSectionResizeMode(QHeaderView.Interactive)
        self.diff_table.horizontalHeader().setStretchLastSection(True)
        for col, width in enumerate((200, 135, 160, 160, 220, 150)):
            self.diff_table.setColumnWidth(col, width)
        self.diff_table.verticalHeader().setVisible(False)
        self.diff_table.verticalHeader().setDefaultSectionSize(36)
        self.diff_table.setAlternatingRowColors(True)
        self.diff_table.setSelectionBehavior(QTableWidget.SelectRows)
        self.diff_table.setEditTriggers(QTableWidget.NoEditTriggers)
        layout.addWidget(self.diff_table, 1)
        self.diff_detail, self.diff_preview = self._add_detail_panel(layout, "Dòng đang có trong file theo dõi")
        self.diff_table.itemSelectionChanged.connect(self.update_diff_context)
        actions = QHBoxLayout()
        self.scope = label("", "hint")
        actions.addWidget(self.scope)
        self.diff_defer = button("Để chứng từ này xử lý sau", lambda: self.defer_selected("diff"))
        self.diff_defer.setProperty("quiet", "true")
        actions.addWidget(self.diff_defer)
        actions.addStretch()
        self.field_apply = button("Lưu cách xử lý số liệu", self.apply_fields, True)
        actions.addWidget(self.field_apply)
        layout.addLayout(actions)
        return page

    def _add_detail_panel(self, layout: QVBoxLayout, heading: str) -> tuple[QLabel, QTableWidget]:
        layout.addWidget(label(heading, "heading"))
        details = label("Chọn một dòng trong bảng phía trên để xem dữ liệu.", "detail")
        layout.addWidget(details)
        preview = QTableWidget(0, 0)
        preview.setMinimumHeight(150)
        preview.setMaximumHeight(225)
        preview.verticalHeader().setVisible(False)
        preview.verticalHeader().setDefaultSectionSize(35)
        preview.setAlternatingRowColors(True)
        preview.setEditTriggers(QTableWidget.NoEditTriggers)
        preview.setSelectionBehavior(QTableWidget.SelectRows)
        preview.horizontalHeader().setSectionResizeMode(QHeaderView.Interactive)
        layout.addWidget(preview)
        return details, preview

    def _review_page(self):
        page, layout = self._page(
            "Xem trước kết quả nhập",
            "Số dòng dưới đây do bộ phân tích PM tính từ lựa chọn hiện tại. Demo chỉ dừng ở bước xem trước, không ghi Excel.",
        )
        self.review_summary = label("", "notice")
        layout.addWidget(self.review_summary)
        self.review_attention = label("", "detail")
        layout.addWidget(self.review_attention)
        self.review_table = QTableWidget(0, 5)
        self.review_table.setHorizontalHeaderLabels(["Phần dữ liệu", "Việc sẽ làm", "Số quyết toán", "Container", "Nơi ghi trong file"])
        self.review_table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.review_table.setEditTriggers(QTableWidget.NoEditTriggers)
        layout.addWidget(self.review_table, 1)
        actions = QHBoxLayout()
        actions.addStretch()
        self.finish_button = button("Mô phỏng xác nhận nhập", self.finish, True)
        actions.addWidget(self.finish_button)
        layout.addLayout(actions)
        return page

    def current_issues(self, kinds: set[str]) -> list[ImportIssue]:
        return [issue for issue in self.sandbox.analysis.issues if issue.kind in kinds]

    def source_row(self, issue: ImportIssue) -> dict[str, Any]:
        return self.sandbox.payload["du_lieu_boc_tach"][issue.json_indexes[0]]

    def source_label(self, issue: ImportIssue) -> str:
        row = self.source_row(issue)
        reference = issue.reference or row.get("file_nguon") or issue.doc_label
        return f"{issue.doc_label} · {reference}"

    def effective_day_key(self, issue: ImportIssue) -> tuple[str, str]:
        row = self.source_row(issue)
        if issue.doc_type == DOC_EXPENSE and parse_date(row.get("ngay_chay")):
            return "ngay_chay", "Ngày chạy"
        return "ngay_dong", "Ngày đóng"

    def _source_summary(self, issue: ImportIssue) -> str:
        row = self.source_row(issue)
        fields = (
            [("File", "file_nguon"), ("Ngày chạy", "ngay_chay"), ("Ngày đóng", "ngay_dong"),
             ("Container", "so_container"), ("Số hóa đơn", "so_hd"), ("Đơn giá", "don_gia"),
             ("Nơi đóng", "noi_dong")]
            if issue.doc_type == DOC_EXPENSE else
            [("File", "file_nguon"), ("Ngày đóng", "ngay_dong"), ("Container", "so_container"),
             ("Số Seal", "so_chi_seal"), ("Số tấn", "so_tan"), ("Nơi đóng", "noi_dong")]
        )
        parts = [f"<b>{html.escape(name)}:</b> {html.escape(shown(row.get(key)) or 'Chưa có')}" for name, key in fields]
        return f"<b>Chứng từ đang xử lý · {html.escape(issue.doc_label)}</b><br>" + " &nbsp; · &nbsp; ".join(parts)

    @staticmethod
    def _set_preview(table: QTableWidget, headers: list[str], rows: list[list[Any]], chosen: int = -1):
        table.setUpdatesEnabled(False)
        table.clear()
        table.setColumnCount(len(headers))
        table.setHorizontalHeaderLabels(headers)
        table.setRowCount(len(rows))
        for row_index, values in enumerate(rows):
            for col_index, value in enumerate(values):
                cell = read_only(shown(value))
                cell.setToolTip(str(value) if value not in (None, "") else "")
                if row_index == chosen:
                    cell.setBackground(QColor("#E5F1FF"))
                table.setItem(row_index, col_index, cell)
        table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeToContents)
        table.horizontalHeader().setStretchLastSection(True)
        table.setUpdatesEnabled(True)

    def _matching_workbook_rows(self, container: str, *, expense: bool = False):
        for (sheet, row_number), values in self.sandbox.workbook_rows.items():
            if (sheet == EXPENSE_SHEET) != expense:
                continue
            if str(values.get("Số Container") or "").strip().upper() == container.strip().upper():
                yield sheet, row_number, values

    def update_edit_context(self):
        if self._refreshing:
            return
        self.edit_defer.setEnabled(0 <= self.edit_table.currentRow() < len(self.edit_ids))
        index = self.edit_table.currentRow()
        if not 0 <= index < len(self.edit_ids):
            self.edit_detail.setText("Chọn một dòng cần sửa để xem chứng từ và các dòng cùng Container.")
            self._set_preview(self.edit_preview, ["Dòng file theo dõi"], [])
            return
        issue = next(item for item in self.sandbox.analysis.issues if item.issue_id == self.edit_ids[index])
        source = self.source_row(issue)
        key, key_label = self.effective_day_key(issue)
        edited_day = self.edit_table.item(index, 3).text().strip()
        edited_container = self.edit_table.item(index, 4).text().strip().upper()
        draft = ""
        if (edited_day, edited_container) != (shown(source.get(key)), str(source.get("so_container") or "").upper()):
            draft = (f"<br><b>Đang sửa, chưa kiểm tra:</b> {html.escape(key_label)} → "
                     f"{html.escape(edited_day or 'Chưa có')} · Container → {html.escape(edited_container or 'Chưa có')}")
        self.edit_detail.setText(self._source_summary(issue) + draft +
                                 "<br><b>Gợi ý đối chiếu:</b> Các dòng dưới đây cùng Container; hãy kiểm tra ngày trước khi sửa.")
        rows = [
            [values.get("SQT PM"), values.get("Ngày Đóng"), values.get("Số Container"),
             values.get("Nơi đóng"), values.get("Loại hàng"), values.get("Người nhận"), f"{sheet}, dòng {number}"]
            for sheet, number, values in self._matching_workbook_rows(issue.container)
        ]
        self._set_preview(self.edit_preview,
                          ["Số quyết toán", "Ngày đóng", "Container", "Nơi đóng", "Loại hàng", "Người nhận", "Vị trí trong file"], rows)
        if not rows:
            self.edit_detail.setText(self.edit_detail.text() + "<br>Chưa có dòng cùng Container trong file theo dõi.")

    def update_choice_context(self):
        if self._refreshing:
            return
        self.choice_defer.setEnabled(0 <= self.choice_table.currentRow() < len(self.choice_ids))
        index = self.choice_table.currentRow()
        if not 0 <= index < len(self.choice_ids):
            self.choice_detail.setText("Chọn một chứng từ phía trên để so sánh các phương án.")
            self._set_preview(self.choice_preview, ["Phương án"], [])
            return
        issue = next(item for item in self.sandbox.analysis.issues if item.issue_id == self.choice_ids[index])
        self.choice_detail.setText(self._source_summary(issue))
        selected_id = str(self.choice_boxes[index].currentData() or "")
        if issue.kind == ISSUE_MULTIPLE_BILLS:
            headers = ["Phương án", "Số Bill", "Ngày chạy", "Tên tàu", "VT biển", "Container", "File Bill"]
            rows = [
                [f"Bill {number}", candidate.values.get("Số B/L"), candidate.values.get("Ngày chạy"),
                 candidate.values.get("Tên tàu"), candidate.values.get("VT biển"),
                 candidate.values.get("Container"), candidate.values.get("File nguồn")]
                for number, candidate in enumerate(issue.candidates, 1)
            ]
        else:
            headers = ["Phương án", "Số quyết toán", "Ngày đóng", "Container", "Nơi đóng", "Loại hàng", "Người nhận", "Vị trí trong file"]
            rows = []
            for number, candidate in enumerate(issue.candidates, 1):
                values = self.sandbox.workbook_rows.get((candidate.sheet_name, candidate.row_number), candidate.values)
                rows.append([f"Dòng {number}", candidate.sqt, values.get("Ngày Đóng") or candidate.values.get("Ngày đóng"),
                             values.get("Số Container") or candidate.values.get("Container"), values.get("Nơi đóng"),
                             values.get("Loại hàng"), values.get("Người nhận"),
                             f"{candidate.sheet_name}, dòng {candidate.row_number}"])
        chosen = next((i for i, item in enumerate(issue.candidates) if item.candidate_id == selected_id), -1)
        self._set_preview(self.choice_preview, headers, rows, chosen)

    def choose_from_preview(self, row: int, _column: int):
        issue_row = self.choice_table.currentRow()
        if 0 <= issue_row < len(self.choice_boxes) and 0 <= row + 1 < self.choice_boxes[issue_row].count():
            self.choice_boxes[issue_row].setCurrentIndex(row + 1)
            self.update_choice_context()

    def update_diff_context(self):
        if self._refreshing:
            return
        self.diff_defer.setEnabled(0 <= self.diff_table.currentRow() < len(self.diff_rows))
        index = self.diff_table.currentRow()
        if not 0 <= index < len(self.diff_rows):
            self.diff_detail.setText("Chọn một trường phía trên để xem chứng từ và dòng đang có trong file theo dõi.")
            self._set_preview(self.diff_preview, ["Dòng file theo dõi"], [])
            return
        issue_id, conflict = self.diff_rows[index]
        issue = next(item for item in self.sandbox.analysis.issues if item.issue_id == issue_id)
        self.diff_detail.setText(self._source_summary(issue) +
                                 f"<br><b>Đang so sánh:</b> {html.escape(conflict.field_name)} · "
                                 f"File theo dõi: {html.escape(shown(conflict.current_value))} · "
                                 f"Chứng từ mới: {html.escape(shown(conflict.new_value))}")
        values = self.sandbox.workbook_rows.get((conflict.target_sheet, conflict.target_row), {})
        if issue.doc_type == DOC_EXPENSE:
            headers = ["Số quyết toán", "Ngày tháng", "Container", "Số hóa đơn", "Nơi đóng", "Đơn giá", "Tổng tiền", "Vị trí trong file"]
            row = [values.get("SQT PM"), values.get("Ngày tháng"), values.get("Số Container"),
                   values.get("Số HĐ"), values.get("Nơi đóng"), values.get("Đơn giá"),
                   values.get("Tổng tiền"), f"Bảng Khoản chi, dòng {conflict.target_row}"]
        else:
            headers = ["Số quyết toán", "Ngày đóng", "Container", "Số Bill", "Số Seal", "Số tấn", "Nơi đóng", "Vị trí trong file"]
            row = [values.get("SQT PM"), values.get("Ngày Đóng"), values.get("Số Container"),
                   values.get("Số B/L"), values.get("Số chì/Seal"), values.get("Số tấn"),
                   values.get("Nơi đóng"), f"{conflict.target_sheet}, dòng {conflict.target_row}"]
        self._set_preview(self.diff_preview, headers, [row] if values else [])

    @staticmethod
    def _set_status(table: QTableWidget, row: int, column: int, value: str):
        item = table.item(row, column)
        if item and item.text() != value:
            blocked = table.blockSignals(True)
            item.setText(value)
            table.blockSignals(blocked)

    def update_edit_actions(self, *_args):
        if self._refreshing:
            return
        changed = 0
        for index, issue_id in enumerate(self.edit_ids):
            issue = next(item for item in self.sandbox.analysis.issues if item.issue_id == issue_id)
            source = self.source_row(issue)
            key, _ = self.effective_day_key(issue)
            day_item, container_item = self.edit_table.item(index, 3), self.edit_table.item(index, 4)
            if day_item is None or container_item is None:
                continue
            differs = (day_item.text().strip(), container_item.text().strip().upper()) != (
                shown(source.get(key)), str(source.get("so_container") or "").upper())
            changed += differs
            self._set_status(self.edit_table, index, 5, "Chưa kiểm tra" if differs else "Cần sửa")
        self.edit_apply.setEnabled(changed > 0)
        self.edit_apply.setText(f"Kiểm tra lại {changed} dòng đã sửa" if changed else "Kiểm tra lại dòng đã sửa")
        self.update_edit_context()

    def update_choice_actions(self):
        if self._refreshing:
            return
        changed = 0
        for index, (issue_id, box) in enumerate(zip(self.choice_ids, self.choice_boxes)):
            candidate_id = str(box.currentData() or "")
            saved = self.sandbox.analysis.resolutions.get(issue_id)
            is_saved = bool(saved and saved.action == RESOLUTION_SELECT and saved.candidate_id == candidate_id)
            pending = bool(candidate_id and not is_saved)
            changed += pending
            status = ("Chưa áp dụng" if pending else "Đã áp dụng" if is_saved else
                      "Để sau" if saved and saved.action == RESOLUTION_SKIP else "Chưa chọn")
            self._set_status(self.choice_table, index, 5, status)
        self.choice_apply.setEnabled(changed > 0)
        self.choice_apply.setText(f"Ghép {changed} lựa chọn" if changed else "Ghép các lựa chọn")

    def _current_field_choices(self, issue: ImportIssue) -> dict[str, dict[str, str]]:
        choices = {}
        for row, (issue_id, conflict) in enumerate(self.diff_rows):
            if issue_id != issue.issue_id:
                continue
            mode = str(self.diff_boxes[row].currentData() or "")
            value = self.diff_inputs[row].text().strip()
            if mode and (mode != "CUSTOM" or self._valid_custom(conflict, value)):
                choices[conflict.conflict_id] = {"mode": mode, "value": value}
        return choices

    def update_field_actions(self, *_args):
        if self._refreshing:
            return
        ready = 0
        for issue in self.current_issues({ISSUE_FIELD_CONFLICT}):
            saved = self.sandbox.analysis.resolutions.get(issue.issue_id)
            if saved and saved.action == RESOLUTION_FIELDS:
                for row, (issue_id, conflict) in enumerate(self.diff_rows):
                    if issue_id == issue.issue_id and not self.diff_boxes[row].currentData():
                        prior_mode = saved.field_choices.get(conflict.conflict_id, {}).get("mode")
                        prior_index = self.diff_boxes[row].findData(prior_mode)
                        if prior_index > 0:
                            blocked = self.diff_boxes[row].blockSignals(True)
                            self.diff_boxes[row].setCurrentIndex(prior_index)
                            self.diff_boxes[row].blockSignals(blocked)
                            self.diff_inputs[row].setEnabled(prior_mode == "CUSTOM")
            choices = self._current_field_choices(issue)
            complete = len(choices) == len(issue.conflicts)
            is_saved = bool(saved and saved.action == RESOLUTION_FIELDS and saved.field_choices == choices)
            ready += bool(complete and not is_saved)
            for row, (issue_id, conflict) in enumerate(self.diff_rows):
                if issue_id != issue.issue_id:
                    continue
                mode = str(self.diff_boxes[row].currentData() or "")
                status = ("Chưa áp dụng" if complete and not is_saved else
                          "Đã áp dụng" if is_saved else
                          "Cần nhập giá trị" if mode == "CUSTOM" and conflict.conflict_id not in choices else
                          "Để sau" if saved and saved.action == RESOLUTION_SKIP and not mode else
                          "Chưa chọn")
                self._set_status(self.diff_table, row, 6, status)
        self.field_apply.setEnabled(ready > 0)
        self.field_apply.setText(f"Lưu cách xử lý cho {ready} chứng từ" if ready else "Lưu cách xử lý số liệu")
        self.update_rule_action()

    def update_rule_action(self, *_args):
        if self._refreshing:
            return
        fillable = sum(
            not self.diff_table.isRowHidden(row) and not box.currentData()
            for row, box in enumerate(self.diff_boxes)
        )
        self.rule_fill.setEnabled(bool(self.rule.currentData()) and fillable > 0)
        self.rule_fill.setText(f"Điền cho {fillable} trường đang hiện" if fillable else "Điền cho các dòng đang hiện")

    def _advance_if_done(self, step: int):
        unresolved = self.sandbox.analysis.unresolved_issues
        for next_step, kinds in ((0, EDIT_KINDS), (1, CHOOSE_KINDS), (2, {ISSUE_FIELD_CONFLICT})):
            if next_step > step and any(issue.kind in kinds for issue in unresolved):
                self.tabs.setCurrentIndex(next_step)
                return
        if not unresolved:
            self.tabs.setCurrentIndex(3)

    def capture_drafts(self):
        for row, issue_id in enumerate(self.edit_ids):
            day = self.edit_table.item(row, 3)
            container = self.edit_table.item(row, 4)
            if day and container:
                self.edit_drafts[issue_id] = (day.text(), container.text())
        for issue_id, box in zip(self.choice_ids, self.choice_boxes):
            self.choice_drafts[issue_id] = str(box.currentData() or "")
        for (_issue_id, conflict), box, editor in zip(self.diff_rows, self.diff_boxes, self.diff_inputs):
            self.field_drafts[conflict.conflict_id] = (str(box.currentData() or ""), editor.text())

    def refresh_ui(self, *, capture: bool = True):
        if capture:
            self.capture_drafts()
        self._refreshing = True
        analysis = self.sandbox.analysis
        pending = analysis.unresolved_issues
        edit_count = sum(issue.kind in EDIT_KINDS for issue in pending)
        choice_count = sum(issue.kind in CHOOSE_KINDS for issue in pending)
        field_count = sum(issue.kind == ISSUE_FIELD_CONFLICT for issue in pending)
        self.summary.setText(
            f"Còn {len(pending)} việc cần quyết định  ·  "
            f"{edit_count} cần sửa thông tin  ·  {choice_count} cần chọn Bill/số quyết toán  ·  "
            f"{field_count} cần chọn số liệu  ·  {analysis.skipped_count} dòng để sau"
        )
        self.tabs.setTabText(0, f"1. Sửa thông tin ({edit_count})")
        self.tabs.setTabText(1, f"2. Chọn Bill / số quyết toán ({choice_count})")
        self.tabs.setTabText(2, f"3. Chọn số liệu ({field_count})")
        self._fill_edit_table()
        self._fill_choice_table()
        self._fill_diff_table()
        for table in (self.edit_table, self.choice_table, self.diff_table):
            if table.rowCount() and table.currentRow() < 0:
                table.selectRow(0)
        self._refreshing = False
        self.update_edit_context()
        self.update_choice_context()
        self.update_diff_context()
        self.update_edit_actions()
        self.update_choice_actions()
        self.update_field_actions()
        self.refresh_review()

    def _fill_edit_table(self):
        issues = [issue for issue in self.current_issues(EDIT_KINDS) if issue in self.sandbox.analysis.unresolved_issues]
        self.edit_ids = [issue.issue_id for issue in issues]
        self.edit_table.setRowCount(len(issues))
        for index, issue in enumerate(issues):
            row = self.source_row(issue)
            key, key_label = self.effective_day_key(issue)
            day, container = self.edit_drafts.get(issue.issue_id, (shown(row.get(key)), str(row.get("so_container") or "")))
            for col, value in enumerate((self.source_label(issue), user_issue_name(issue), key_label)):
                self.edit_table.setItem(index, col, read_only(value))
            self.edit_table.setItem(index, 3, QTableWidgetItem(day))
            self.edit_table.setItem(index, 4, QTableWidgetItem(container))
            self.edit_table.setItem(index, 5, read_only("Cần sửa"))

    def _fill_choice_table(self):
        issues = self.current_issues(CHOOSE_KINDS)
        self.choice_ids = [issue.issue_id for issue in issues]
        self.choice_boxes = []
        self.choice_table.setRowCount(len(issues))
        for index, issue in enumerate(issues):
            for col, value in enumerate((self.source_label(issue), user_issue_name(issue), shown(issue.match_date), issue.container)):
                self.choice_table.setItem(index, col, read_only(value))
            box = QComboBox()
            box.addItem("Chọn cho dòng này…", "")
            for candidate in issue.candidates:
                box.addItem(user_candidate_name(candidate), candidate.candidate_id)
            saved = self.sandbox.analysis.resolutions.get(issue.issue_id)
            wanted = self.choice_drafts.get(issue.issue_id) or (saved.candidate_id if saved and saved.action == RESOLUTION_SELECT else "")
            index_in_box = box.findData(wanted)
            if index_in_box >= 0:
                box.setCurrentIndex(index_in_box)
            box.currentIndexChanged.connect(lambda _value, row=index: self._choice_changed(row))
            self.choice_table.setCellWidget(index, 4, box)
            self.choice_boxes.append(box)
            self.choice_table.setItem(index, 5, read_only("Để sau" if saved and saved.action == RESOLUTION_SKIP else "Đã chọn" if saved else "Chưa chọn"))
    def _choice_changed(self, row: int):
        if self._refreshing:
            return
        issue_id = self.choice_ids[row]
        saved = self.sandbox.analysis.resolutions.get(issue_id)
        if saved and saved.action == RESOLUTION_SELECT and not self.choice_boxes[row].currentData():
            prior_index = self.choice_boxes[row].findData(saved.candidate_id)
            if prior_index > 0:
                blocked = self.choice_boxes[row].blockSignals(True)
                self.choice_boxes[row].setCurrentIndex(prior_index)
                self.choice_boxes[row].blockSignals(blocked)
        self.choice_table.selectRow(row)
        self.update_choice_context()
        self.update_choice_actions()

    def _fill_diff_table(self):
        pairs = [(issue, conflict) for issue in self.current_issues({ISSUE_FIELD_CONFLICT}) for conflict in issue.conflicts]
        self.diff_rows = [(issue.issue_id, conflict) for issue, conflict in pairs]
        self.diff_boxes = []
        self.diff_inputs = []
        self.diff_table.setRowCount(len(pairs))
        for row_index, (issue, conflict) in enumerate(pairs):
            values = (self.source_label(issue), conflict.field_name, shown(conflict.current_value), shown(conflict.new_value))
            for col, value in enumerate(values):
                self.diff_table.setItem(row_index, col, read_only(value))
            box = QComboBox()
            box.addItem("Chưa chọn", "")
            box.addItem("Giữ số liệu đang có", "CURRENT")
            box.addItem("Dùng số liệu mới", "NEW")
            box.addItem("Nhập giá trị khác", "CUSTOM")
            editor = QLineEdit()
            saved = self.sandbox.analysis.resolutions.get(issue.issue_id)
            saved_choice = saved.field_choices.get(conflict.conflict_id, {}) if saved else {}
            mode, value = self.field_drafts.get(conflict.conflict_id, (str(saved_choice.get("mode") or ""), shown(saved_choice.get("value"))))
            box.setCurrentIndex(max(0, box.findData(mode)))
            editor.setText(value)
            editor.setEnabled(box.currentData() == "CUSTOM")
            box.currentIndexChanged.connect(lambda _index, target=editor, source=box: target.setEnabled(source.currentData() == "CUSTOM"))
            box.currentIndexChanged.connect(self.update_field_actions)
            editor.textChanged.connect(self.update_field_actions)
            self.diff_table.setCellWidget(row_index, 4, box)
            self.diff_table.setCellWidget(row_index, 5, editor)
            self.diff_boxes.append(box)
            self.diff_inputs.append(editor)
            self.diff_table.setItem(row_index, 6, read_only("Để sau" if saved and saved.action == RESOLUTION_SKIP else "Đã chọn" if saved else "Chưa chọn"))
        old_field = self.field_filter.currentText()
        self.field_filter.blockSignals(True)
        self.field_filter.clear()
        self.field_filter.addItem("Tất cả")
        self.field_filter.addItems(sorted({conflict.field_name for _, conflict in pairs}))
        self.field_filter.setCurrentIndex(max(0, self.field_filter.findText(old_field)))
        self.field_filter.blockSignals(False)
        self.filter_fields()

    def filter_fields(self):
        wanted = self.field_filter.currentText()
        visible = 0
        for row, (_issue_id, conflict) in enumerate(self.diff_rows):
            show = wanted == "Tất cả" or wanted == conflict.field_name
            self.diff_table.setRowHidden(row, not show)
            visible += show
        self.scope.setText(f"Đang hiện {visible}/{len(self.diff_rows)} trường")
        self.update_rule_action()

    def fill_edit_example(self):
        for index, issue_id in enumerate(self.edit_ids):
            issue = next(issue for issue in self.sandbox.analysis.issues if issue.issue_id == issue_id)
            if issue.kind == ISSUE_MISSING_CLOSE_DATE:
                self.edit_table.item(index, 3).setText("16/06/2026")
            elif issue.kind == ISSUE_NO_TARGET:
                self.edit_table.item(index, 3).setText("15/06/2026")
        self.feedback.setText("Đã điền ngày mẫu. Bấm ‘Kiểm tra lại các dòng đã sửa’ để PM ghép lại.")

    def fill_choice_example(self):
        for box in self.choice_boxes:
            if box.count() > 1:
                box.setCurrentIndex(1)
        self.update_choice_actions()
        self.feedback.setText("Đã chọn ứng viên đầu tiên cho mỗi dòng. Bạn có thể đổi riêng từng dòng trước khi áp dụng.")

    def fill_rule(self):
        mode = str(self.rule.currentData() or "")
        if not mode:
            QMessageBox.information(self, "Chưa chọn cách xử lý", "Hãy chọn cách xử lý trước khi điền vào bảng.")
            return
        count = 0
        for row, box in enumerate(self.diff_boxes):
            if not self.diff_table.isRowHidden(row) and not box.currentData():
                box.setCurrentIndex(box.findData(mode))
                count += 1
        self.update_field_actions()
        self.feedback.setText(f"Đã điền cách xử lý cho {count} trường đang hiện. Kiểm tra ngoại lệ rồi bấm Lưu cách xử lý.")

    def apply_edits(self):
        changes = []
        for index, issue_id in enumerate(self.edit_ids):
            issue = next(issue for issue in self.sandbox.analysis.issues if issue.issue_id == issue_id)
            row = self.source_row(issue)
            key, _label = self.effective_day_key(issue)
            day = self.edit_table.item(index, 3).text().strip()
            container = self.edit_table.item(index, 4).text().strip().upper()
            if (day, container) == (shown(row.get(key)), str(row.get("so_container") or "").upper()):
                continue
            if day and parse_date(day) is None:
                QMessageBox.warning(self, "Ngày chưa đúng", f"Dòng {index + 1}: hãy nhập ngày theo dạng 15/06/2026.")
                return
            changes.append((issue.json_indexes[0], key, day, container, issue_id))
        if not changes:
            self.feedback.setText("Chưa có ngày hoặc Container nào thay đổi.")
            return
        self.capture_drafts()
        for json_index, key, day, container, issue_id in changes:
            row = self.sandbox.payload["du_lieu_boc_tach"][json_index]
            row[key] = day
            row["so_container"] = container
            self.edit_drafts.pop(issue_id, None)
        old_ids = {issue.issue_id for issue in self.sandbox.analysis.unresolved_issues}
        self.sandbox.reanalyze()
        new_ids = {issue.issue_id for issue in self.sandbox.analysis.unresolved_issues}
        self.last_deferred = None
        self.undo_button.setVisible(False)
        self.refresh_ui(capture=False)
        self.feedback.setText(f"PM đã kiểm tra lại {len(changes)} dòng: {len(old_ids - new_ids)} vấn đề hết, {len(new_ids - old_ids)} vấn đề mới phát sinh.")
        self._advance_if_done(0)

    def apply_choices(self):
        selected = []
        for issue_id, box in zip(self.choice_ids, self.choice_boxes):
            candidate_id = str(box.currentData() or "")
            saved = self.sandbox.analysis.resolutions.get(issue_id)
            if candidate_id and (saved is None or saved.action != RESOLUTION_SELECT or saved.candidate_id != candidate_id):
                selected.append((issue_id, candidate_id))
        if not selected:
            self.feedback.setText("Chưa có lựa chọn mới nào cần áp dụng.")
            return
        self.capture_drafts()
        for issue_id, candidate_id in selected:
            self.sandbox.resolutions[issue_id] = IssueResolution(issue_id, RESOLUTION_SELECT, candidate_id=candidate_id)
        old_ids = {issue.issue_id for issue in self.sandbox.analysis.unresolved_issues}
        self.sandbox.reanalyze()
        new_ids = {issue.issue_id for issue in self.sandbox.analysis.unresolved_issues}
        self.last_deferred = None
        self.undo_button.setVisible(False)
        self.refresh_ui(capture=False)
        self.feedback.setText(f"PM đã ghép lại {len(selected)} chứng từ: {len(old_ids - new_ids)} vấn đề hết, {len(new_ids - old_ids)} vấn đề mới phát sinh. Bill còn lại giữ PENDING để xử lý sau.")
        self._advance_if_done(1)

    def _valid_custom(self, conflict, value: str) -> bool:
        reference = conflict.new_value if conflict.new_value not in (None, "") else conflict.current_value
        if isinstance(reference, (date, datetime)):
            return parse_date(value) is not None
        if isinstance(reference, (int, float)) and not isinstance(reference, bool):
            return parse_number(value) is not None
        return bool(value)

    def apply_fields(self):
        self.capture_drafts()
        to_apply = []
        for issue in self.current_issues({ISSUE_FIELD_CONFLICT}):
            choices = {}
            for conflict in issue.conflicts:
                mode, value = self.field_drafts.get(conflict.conflict_id, ("", ""))
                if mode == "CUSTOM" and not self._valid_custom(conflict, value):
                    QMessageBox.warning(self, "Giá trị chưa đúng", f"Hãy kiểm tra giá trị nhập cho trường ‘{conflict.field_name}’.")
                    return
                if mode:
                    choices[conflict.conflict_id] = {"mode": mode, "value": value}
            if len(choices) == len(issue.conflicts):
                saved = self.sandbox.analysis.resolutions.get(issue.issue_id)
                if saved is None or saved.action != RESOLUTION_FIELDS or saved.field_choices != choices:
                    to_apply.append((issue.issue_id, choices))
        if not to_apply:
            self.feedback.setText("Chọn đủ cách xử lý cho mọi trường của một chứng từ để áp dụng.")
            return
        for issue_id, choices in to_apply:
            self.sandbox.resolutions[issue_id] = IssueResolution(issue_id, RESOLUTION_FIELDS, field_choices=choices)
        self.sandbox.reanalyze()
        self.last_deferred = None
        self.undo_button.setVisible(False)
        self.refresh_ui(capture=False)
        self.feedback.setText(f"PM đã ghi nhận lựa chọn số liệu cho {len(to_apply)} chứng từ. File theo dõi mẫu vẫn chưa được ghi.")
        self._advance_if_done(2)

    def defer_selected(self, group: str):
        table, ids = {"edit": (self.edit_table, self.edit_ids),
                      "choice": (self.choice_table, self.choice_ids),
                      "diff": (self.diff_table, [issue_id for issue_id, _ in self.diff_rows])}[group]
        row = table.currentRow()
        if not 0 <= row < len(ids):
            QMessageBox.information(self, "Chưa chọn dòng", "Hãy chọn một dòng trong bảng.")
            return
        issue_id = ids[row]
        issue = next(issue for issue in self.sandbox.analysis.issues if issue.issue_id == issue_id)
        self.capture_drafts()
        self.last_deferred = (issue_id, self.sandbox.resolutions.get(issue_id))
        for conflict in issue.conflicts:
            self.field_drafts.pop(conflict.conflict_id, None)
        self.choice_drafts.pop(issue_id, None)
        self.sandbox.resolutions[issue_id] = IssueResolution(issue_id, RESOLUTION_SKIP)
        self.sandbox.reanalyze()
        self.refresh_ui(capture=False)
        self.undo_button.setVisible(True)
        self.feedback.setText(f"Đã để {issue.doc_label.lower()} này sang lần nhập sau. Có thể hoàn tác ngay.")
        self._advance_if_done({"edit": 0, "choice": 1, "diff": 2}[group])

    def undo_defer(self):
        if not self.last_deferred:
            return
        issue_id, previous = self.last_deferred
        if previous is None:
            self.sandbox.resolutions.pop(issue_id, None)
        else:
            self.sandbox.resolutions[issue_id] = previous
        self.last_deferred = None
        self.undo_button.setVisible(False)
        self.sandbox.reanalyze()
        self.refresh_ui(capture=False)
        self.feedback.setText("Đã hoàn tác. Chứng từ trở lại trạng thái trước khi để sau.")

    def on_tab_change(self, index: int):
        if index == 3:
            self.refresh_review()

    def refresh_review(self):
        analysis = self.sandbox.analysis
        waiting = len(analysis.unresolved_issues)
        self.review_summary.setText(
            f"Dự kiến tạo {analysis.new_info_count} dòng quyết toán, cập nhật {analysis.updated_info_count} dòng quyết toán, "
            f"ghi {len(analysis.expense_changes)} dòng khoản chi; {analysis.skipped_count} dòng để sau. "
            + (f"Còn {waiting} việc cần quyết định trước khi xác nhận." if waiting else "Đã đủ quyết định để xác nhận.")
        )
        waiting_names = [user_issue_name(issue) for issue in analysis.unresolved_issues]
        skipped_names = []
        for index in sorted(set(analysis.skipped_json_indexes)):
            if 0 <= index < len(self.sandbox.payload["du_lieu_boc_tach"]):
                source = self.sandbox.payload["du_lieu_boc_tach"][index]
                identifier = source.get("so_bill") or source.get("so_hd") or source.get("file_nguon") or "Chứng từ"
                skipped_names.append(f"{source.get('loai_chung_tu') or 'Chứng từ'} {identifier}")
        details = []
        if waiting_names:
            details.append("<b>Cần quyết định:</b> " + html.escape(" · ".join(waiting_names)))
        if skipped_names:
            details.append("<b>Giữ PENDING để xử lý sau:</b> " + html.escape(" · ".join(skipped_names)))
        self.review_attention.setText("<br>".join(details) if details else "<b>Không có chứng từ nào cần xử lý sau.</b>")
        changes = []
        for change in analysis.info_changes:
            changes.append(("Quyết toán", "Tạo dòng mới" if change.action == "CREATE" else "Cập nhật dòng", change.sqt,
                            change.values.get("Số Container"), change.target_sheet or "Thông tin quyết toán"))
        for change in analysis.expense_changes:
            changes.append(("Khoản chi", "Tạo dòng mới" if change.action == "CREATE" else "Cập nhật dòng", change.values.get("SQT PM"),
                            change.values.get("Số Container"), "Bảng Khoản chi"))
        self.review_table.setRowCount(len(changes))
        for row, values in enumerate(changes):
            for col, value in enumerate(values):
                self.review_table.setItem(row, col, read_only(value))
        self.finish_button.setEnabled(waiting == 0)

    def finish(self):
        QMessageBox.information(self, "Xác nhận nhập · Demo", "Đây là bước xác nhận trong demo. Không file theo dõi thật nào được ghi.")

    def reset(self):
        self.last_deferred = None
        self.undo_button.setVisible(False)
        self.edit_drafts.clear()
        self.choice_drafts.clear()
        self.field_drafts.clear()
        self.sandbox.reset()
        self.tabs.setCurrentIndex(0)
        self.refresh_ui(capture=False)
        self.feedback.setText("Đã trở lại 7 tình huống ban đầu; PM phân tích lại từ dữ liệu mẫu.")

    def closeEvent(self, event):  # noqa: N802 - Qt API
        self.sandbox.close()
        super().closeEvent(event)


def main() -> int:
    app = QApplication(sys.argv)
    font_file = Path("C:/Windows/Fonts/segoeui.ttf")
    if font_file.is_file():
        QFontDatabase.addApplicationFont(str(font_file))
    app.setStyleSheet(STYLE)
    window = StandardBulkDemo()
    window.show()
    return app.exec()


if __name__ == "__main__":
    raise SystemExit(main())
